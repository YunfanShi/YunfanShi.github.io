from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
root=Path('../docs/backend-load-and-errors-3.25.0');root.mkdir(parents=True,exist_ok=True)
d=Document();sec=d.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Microsoft YaHei';st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.2;d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(15)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
d.add_heading('JackYun 云端减负与登录信任改造说明',0)
p('网站 3.25.0    Azure 后台 3.6.0    2026 年 10 月 17 日')
p('本次根据现有生产配置、源码和只读数据库证据，优先消除重复大包读取与远程登录核验，并完成关闭同类错误后静默 30 分钟的功能。主要工作区同步已在 Azure；其余业务尚未全部切换。因此本次采用先减负、再逐模块迁移的方式，保留现有用户数据、账号标识、权限边界与回退能力。')
h('当前架构核对')
t=d.add_table(rows=1,cols=3);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['范围','当前生产归属','依据与边界']):c.text=v
for row in [('工作区与旧页面同步','Azure Functions 与 Azure PostgreSQL','生产 AZURE_SYNC_API_URL 已配置，源端 5 个 ALWAYS 围栏仍在；Azure 健康接口与版本可读取。'),('登录与令牌签发','Supabase Auth','公开 JWKS 为 ES256，可使用本地签名验证；刷新令牌和新登录仍需要签发服务。'),('账号停用与资料','Supabase 为当前权威源','Azure 原有资料并非持续更新的权威镜像，不能直接取消状态同步。'),('通知及其余业务表','Supabase','通知快照、部分阅读、学习、设置等业务仍通过 Supabase 访问。'),('头像与图书文件','生产接口仍使用 Supabase','历史 Azure 复制校验不代表当前新增和删除已经追平。')]:
 for c,v in zip(t.add_row().cells,row):c.text=v
t.autofit=False
for col,width in zip(t.columns,[1.5,1.9,3.5]):col.width=Inches(width)
from docx.oxml import OxmlElement
for row in t.rows:
 for cell,width in zip(row.cells,[1.5,1.9,3.5]):
  cell.width=Inches(width)
  margins=OxmlElement('w:tcMar')
  for edge in ['top','left','bottom','right']:
   node=OxmlElement('w:'+edge);node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');margins.append(node)
  cell._tc.get_or_add_tcPr().append(margins)
  for para in cell.paragraphs:
   para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.15
   for run in para.runs:run.font.size=Pt(10)
h('判断')
p('可以继续把大部分业务移向 Azure，但只改连接地址不会完成迁移。Azure 账号门禁、业务 RPC、RLS 身份上下文、文件授权和新增删除追平必须一起适配。尤其不能把历史恢复库直接作为当前业务权威库，也不能删掉 Azure 入口来回退到已落后的旧同步表。')
d.add_page_break()
h('先解决已量出的重复传输')
p('当前计时会话的 JSON 负载为 127509 字节，其中题型配置为 126031 字节。原计时桥接层每 5 秒读取一次云端控制权；控制器另有每 10 秒的心跳，两类调用均会回传整份负载。以测得的 127772 字节完整响应估算，单设备连续运行一小时、18 次每分钟的状态与心跳读取，约重复传输 138 MB，不包括其他请求、协议开销和压缩。这是模型估算，不能直接归因于账单上的全部超额流量。')
p('新增 practice_timer_status：仅返回当前用户的 ownerId、revision、会话存在性、sessionId 和 savedAt。RLS 与显式 user_id 条件同时限制读取；该查询不插入或锁定会话行。新增 practice_timer_control_compact：复用原有控制权与版本围栏，只在 read 和 claim 时返回完整内容，touch、save 和 release 的 payload 在数据库内部移除，避免先传出 Supabase 再在网站裁剪。')
p('实际同一会话测得：完整响应 127772 字节，状态响应 188 字节，紧凑响应包络约 250 字节。由此削减超过 99.8% 的这部分响应体。首次接续仍读取必要数据；无令牌的本机练习轮询只在远端版本改变时下载内容，不会反复拉取同一配置。真实月度节流和数据库负载下降仍需发布后观察。')
h('关闭错误后的静默规则')
p('用户手动关闭某条错误提示后，该统一错误码及其别名在 30 分钟内不再弹窗。不同路径和模块的相同类型一起静默，已显示的同类卡片一起关闭；其他错误码继续显示。静默截止时间按点击关闭的时刻计算，后续重复错误不会延长；原有 30 秒自动消失不触发静默。')
p('静默状态保存在带 debug 前缀的本机键中，沿用同步排除策略，跨页面、刷新、同源标签页及同源 iframe 生效。存储不可用时仍在本页内存中静默。控制台原始输出、客户端日志与最近诊断仍正常采集；复制另一条错误的诊断也可以看到静默期间的错误。')
d.add_page_break()
h('登录后的信任方式')
p('网站新增 getVerifiedUser，调用 SDK getClaims 对令牌签名与有效期进行验证。当前项目公开签名密钥为 ES256，正常请求复用缓存公钥，避免每次到 Auth 的 user 接口重新取身份。优先接入计时器、通知、离线快照、阅读、设置及请求身份上下文等重复调用路径。未信任未验签 Cookie、裸 getSession 结果或未校验的 JWT 解码；过期和错误角色仍拒绝。')
p('管理员仍读取当前角色与账号状态，业务行权限仍依赖已验证的用户 ID 与数据库 RLS。本次没有删除账号系统、跨用户权限控制或所有安全检查。其他尚未接入 helper 的低频与敏感路径保持原行为；后续应按使用量和权限边界逐项收敛，而非机械替换全部认证代码。')
h('Azure 会话缓存')
p('Azure 对同一用户和 session_id 的已确认 active 结果缓存 5 分钟，每个进程最多 500 条，不保存明文令牌。每次仍先验签并校验令牌期限，然后才允许命中账号状态缓存。正常状态变更可能存在最多 5 分钟缓存窗口；看到明确 401、403、停用、删除或身份不符时立即拒绝并清理缓存。')
p('状态服务出现临时超时、网络错误或服务不可用时，只允许此前确认正常的同一会话沿用缓存；有效期最多到原已确认令牌的到期时刻，不无限顺延，失败复核每 30 秒最多重试一次。新会话、首次核验失败、进程冷启动无缓存、过期令牌和明确拒绝都不使用故障回退。旧 HS256 令牌无法只靠公开密钥本地验签，其现有服务端校验保留。')
p('该改造缓解短暂故障造成的连锁中断，但不等于 Azure 已完全脱离 Supabase。令牌签发、刷新、公钥首次获取、冷启动后的首次账号核验仍有服务依赖。彻底解耦需要让 Azure 成为账号生命周期与停用状态的权威源，再决定是否保留 Supabase Auth，或迁移到其他身份签发服务并保持原用户 UUID 映射。')
h('验证')
p('网站相关回归 59 项、同步回归 77 项、Azure 后台 20 项均通过，各组覆盖范围有交叠。正式构建和 TypeScript 通过，改动文件 ESLint、生成物一致性与密钥扫描通过。真实浏览器在桌面和手机视口验证静默 30 分钟、跨页面与标签页、日志保留、剪贴板降级，以及原有离线计时和跨设备控制契约。数据库验证了元数据大小、无 payload、匿名禁止执行与第二用户隔离；带副作用的契约测试均在事务内回滚。')
d.add_page_break()
h('后续迁移的建议顺序')
p('第一批迁移文件：先清点当前头像、图书和附件的元数据与实际对象，比较 Azure 副本的路径、大小、哈希和删除情况，补齐增量。保留私有容器，通过 Azure 后台签发短期下载授权，验证用户归属、失效链接、缓存和撤销；不能为了减少请求把私有图书容器直接公开。静态公开资源可单独使用缓存策略，不携带会话。')
p('第二批迁移账号状态和通知：定义 Azure 的用户 UUID、active 状态、删除与封禁版本。先完成账号创建、修改、停用和删除的可靠写入与对账，再让后台只查 Azure 权威状态，替代逐次跨云核验。通知迁移要同时覆盖发布、收件范围、已读、删除、分页和未读数；保留现有跨标签页协调与后台暂停，避免迁移后又形成轮询洪峰。')
p('第三批迁移剩余业务表和 RPC：按模块明确唯一写入方，先验证普通角色事务身份、RLS、约束、授权函数和备份恢复，再冻结对应源端写入、排空在途请求、最终追平、核对行数和内容指纹。使用当前 UUID 和原幂等操作编号，验收并发、离线队列、删除、复试、导入导出与回退。禁止以无约束的长期双写代替一致性设计。')
h('发布与回退')
p('本次新增错误静默能力，网站按规则从 3.24.5 升至 3.25.0；Azure 后台升至 3.6.0。数据库只添加兼容函数，原 RPC 保留，因此可先应用迁移再发布网站；若网站回退，旧 RPC 与旧客户端仍可用。新增迁移文件编号与远端已应用历史对齐。Azure 会话缓存改动使用原服务、原数据库和原网络，不增加付费资源。')
p('每批数据切换前保留可实际恢复的最新备份与冻结状态记录。回退先暂停目标写入，反向追平切换后的新增、修改与删除，再恢复源入口；不能只换环境变量。发布后观察至少一个完整使用周期的接口次数、响应体大小、错误率、数据库 CPU 和账单用量，并区分实验估算与线上实测。学生订阅额度也需要持续核对，不把现有信用额度视为永久免费容量。')
h('参考依据')
p('Supabase 官方 Manage Egress usage：建议减少返回字段、降低查询频率，并说明历史用量无法通过优化追溯减少。https://supabase.com/docs/guides/platform/manage-your-usage/egress')
p('Supabase JavaScript getClaims：非对称签名可使用缓存 JWKS 验证；旧对称签名仍可能请求认证服务。https://supabase.com/docs/reference/javascript/auth-getclaims')
p('Microsoft Azure Architecture Center Strangler Fig：通过逐模块替换和验证完成迁移。https://learn.microsoft.com/en-us/azure/architecture/patterns/strangler-fig')
for st in d.styles:
 for node in list(st.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 if st.name.startswith(('Title','Heading','heading')):st.font.color.rgb=RGBColor(0,0,0)
for para in d.paragraphs:
 for run in para.runs:run.font.color.rgb=RGBColor(0,0,0)
d.save(root/'云端减负与登录信任改造说明3.25.0.docx')
print('Created 4 page report')
