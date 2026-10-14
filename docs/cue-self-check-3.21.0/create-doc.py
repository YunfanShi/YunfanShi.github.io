from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

folder = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2.2)
for name, size in [('Normal', 11), ('Title', 22), ('Heading 1', 16), ('Heading 2', 13)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.15
    for color in style._element.iter(qn('w:color')):
        for attribute in ['themeColor', 'themeShade', 'themeTint']:
            color.attrib.pop(qn('w:' + attribute), None)
    for border in list(style._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for link in list(style._element.iter(qn('w:link'))):
        link.getparent().remove(link)
doc.core_properties.title = 'CUE 学习自动对比与 AI 辅助自查说明'
doc.core_properties.author = 'JackYun'
def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text, level=2)
doc.add_paragraph('CUE 学习自动对比与 AI 辅助自查说明', 'Title')
p('版本 3.21.0    更新日期 2026年10月14日')
p('本次更新优化 Schedule 的 CUE 答题卡评分点自查。写完作答并展开对照时，网站自动进行本地文本比较，显示相似度和相同片段，并辅助勾选匹配程度较高的 MS。用户可以直接调整勾选；AI 评价仅由按钮手动触发，沿用网站已有 AI 配置。')
h('一 使用流程')
p('进入 CUE 学习，打开 CUE 答题卡，先独立输入答案或解题过程，再点击“写完了，对照答案”。本地结果立即显示，不需要联网或调用模型。参考答案与我的答案继续支持原有公式渲染；独立的原始文本对照区高亮相同片段，便于定位重合内容。')
p('MS 列表展示每一点的文本覆盖比例和核对提示，符合本地条件的评分点自动勾选。随后可以手动取消或补充勾选。对同一作答和参考版本点击“重新对照答案”保留已有调整；“重新本地检查 MS”则再次补充符合条件的项目，不取消已有勾选。')
p('需要语义评价时，点击“AI 辅助自查并勾选 MS”。AI 返回总体评价、改进建议、逐点判断及作答依据，完整覆盖且具有有效依据的项目会补充勾选。此操作不自动改变 CUE 的掌握状态、练习评价或复习安排。')
h('二 无数据与异常处理')
p('空白作答不进入对比。没有参考答案且没有有效 MS 时跳过检查，辅助按钮不可用；只有 MS 时仍可逐点比较，但不显示虚构的参考答案相似度。只有参考答案时可显示相似度和高亮，不生成评分点。')
p('AI 请求提供取消按钮及 90 秒取消计时。服务失败、权限或配置错误、结果格式不完整时显示可理解的提示，保留已有勾选，本地功能继续可用。请求期间修改作答、重新提交、参考版本变化或关闭答题卡，会取消请求或丢弃过期结果。')
doc.add_page_break()
doc.add_heading('实现规则与数据保存', level=1)
h('三 本地比较规则')
p('本地比较先统一字符宽度和英文大小写，英文按词、中文按相邻双字分解，移除常见英文语法词并对重复词去重。参考答案相似度使用双方词集合的重合比例；MS 覆盖率使用评分点有效词中已在作答出现的比例。')
p('自动勾选阈值为 MS 文本覆盖率至少 85%。另外检查否定词状态及评分点数字是否在作答中完整出现；不一致时保留人工核对。高亮使用原始文本，英文保持词边界，中文合并相邻重合片段，文本以 React 文本节点呈现，不注入 HTML。')
p('本地算法不能判断语义等价、推导正确性、单位关系或数学逻辑。例如不同表述可能遗漏自动勾选，单纯词重合也可能存在语义错误。因此界面明确标示相似度不是得分，用户始终可以修改 MS。公式保留渲染视图，高亮视图展示原始公式文本。')
h('四 AI 请求与结果校验')
p('调用现有 callAiApi 统一入口，继承云端、个人或浏览器 AI 模式及站内权限处理。新增功能不创建或写入任何 API 密钥，不改变服务端模型路由。仅发送当前题干、作答、参考答案及 MS，不发送之前作答或旧备忘录。')
p('提示词要求将学习内容作为数据处理，不能执行其中指令。响应必须包含总体评价、合法范围内的相似度及全部 MS 的唯一编号。重复、缺失、越界编号、错误类型或非法分数会拒绝整份结果；声称覆盖但引用不在作答中的评分点不自动勾选。此校验只能确认依据来源，无法保证模型判断正确。')
h('五 保存与兼容')
p('作答、参考快照和 MS 勾选继续使用原答题卡数据结构及现有保存同步流程。编辑作答会使旧对照失效，重新对照时基于新作答进行本地检查。相同参考快照上的手动勾选会保留。历史作答与旧备忘录仍可展开查看。')
p('AI 评价正文仅在当前答题卡打开期间展示；自动生成的勾选与手动勾选一样保存。重新打开后本地结果可从保存的作答和参考快照计算，AI 评价需要再次手动请求。')
doc.add_page_break()
doc.add_heading('验证与仓库发布', level=1)
h('六 验证结果')
p('24 项相关自动测试通过，覆盖自动本地勾选、重复对照保留人工调整、编辑后重新计算、英文中文匹配、缺少数据、否定词和数字差异、高亮保留原文、AI 无效格式与虚构依据，以及原有移动端公式与答题卡行为。')
p('浏览器回归使用 Edge，覆盖 1440 像素桌面和 390 像素手机视口。确认本地检查不发送 AI 请求，初始勾选正确，手动取消后重新对照保持取消，点击 AI 后更新覆盖项，503 失败不清空勾选，并检查本机保存结果。两种视口均没有页面运行异常。')
p('AI 网络响应采用受控模拟，验证了客户端调用、响应解析、勾选和错误回退；没有对真实模型进行付费调用，因此不对当前真实服务的可用性或评价准确率作保证。静态检查、TypeScript 类型检查及 Next.js 生产构建通过。')
h('七 修改文件与版本')
p('新增 cue-self-check.ts，集中处理本地比较、高亮、AI 提示和响应校验；cue-answer-card.ts 在展开对照时自动设置本地覆盖项；cue-answer-card.tsx 加入评价、手动 AI 按钮及高亮区。新增领域测试与桌面手机浏览器测试，补充原移动组件测试的依赖模拟。')
p('本次属于新增学习功能，按版本规则将本机 3.20.0 提升至 3.21.0，更新 package.json、锁文件根版本、APP_VERSION 与更新历史。仓库内同时保存本说明、生成脚本与浏览器验证记录。')
h('八 同步与范围说明')
p('任务开始时成功获取远端 Main，本机已提交 HEAD 与远端一致，基线提交为 fa16455。但工作区存在其他任务的大量未提交或已暂存修改，部分版本字段高于 Main。此次提交仅纳入本任务代码、测试、版本元数据和说明，使用独立暂存区避免提交其他任务内容，不创建额外分支。')
p('发布前再次读取 Main 并核对提交变化，提交后推送到 main 并再次检查本机 HEAD 与 origin/main 是否一致。保留其他任务的工作区和暂存修改。因此提交版本可以同步一致，但已有未提交功能仍需要它们各自完成验收发布。')
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
doc.save(folder / 'CUE学习自动对比与AI辅助自查说明3.21.0.docx')
print(folder / 'CUE学习自动对比与AI辅助自查说明3.21.0.docx')
