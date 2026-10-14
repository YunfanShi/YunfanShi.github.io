from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7);section.left_margin=section.right_margin=Inches(.75)
for name,size in [('Normal',11),('Title',22),('Heading 1',13)]:
    style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].paragraph_format.space_after=Pt(8);doc.styles['Normal'].paragraph_format.line_spacing=1.12
def p(text):doc.add_paragraph(text)
def h(text):return doc.add_heading(text,1)
doc.add_paragraph('JackYun 错误提示显示规则修复说明',style='Title')
p('版本 3.22.1    2026 年 10 月 14 日')
p('本次调整全站共享错误提示的生命周期：前台提示首次出现后显示 30 秒；网站进入后台或浏览器窗口失去焦点时立即清除提示。后台错误仍记录诊断，不在返回前台时补弹。右上角布局、错误码、复制日志与错误码文档入口继续保留。')
h('显示与关闭规则')
p('每条提示保存独立的到期时间，并由最近到期的一条安排统一定时器。到达 30 秒时删除该条提示。显示期间同一错误码和同一来源再次发生时，仅增加重复次数并更新诊断，不重新计时；其他新错误有自己的 30 秒期限，仍最多显示三条。用户手动关闭后同步重新安排计时。')
p('页面可见且窗口拥有焦点时才允许弹出。切换到其他标签页、最小化或切到其他应用造成页面隐藏或失焦时，立即清除所有可见提示并取消显示定时器。页面恢复可见不会重新显示之前的提示；之后在前台发生的新错误正常提示。')
h('后台日志与兼容范围')
p('错误首先经过脱敏并加入近期诊断，然后才判断是否需要显示。后台静默仅影响弹窗，不影响原有日志捕获、同步队列、版本检查或数据保存。前台新提示的复制日志仍可包含近期后台错误，方便排查。共享脚本覆盖登录、管理员页面、旧 HTML 模块和嵌入页面，原有错误码不变。')
h('实现与测试').paragraph_format.page_break_before=True
p('修改 scripts/error-notices-runtime.js，并由 build-error-system.mjs 重新生成 public/shared/error-notices.js、开发文档和在线错误码说明。编号目录未变化，新增到期时间、定时清理、visibilitychange 和 blur 处理。此次属于小问题修复，应用版本从 3.22.0 更新为 3.22.1，后台无需重新部署。')
p('四项错误系统单元测试通过，覆盖 29999 毫秒仍显示、30000 毫秒隐藏、重复错误不延时、页面隐藏后清空、后台错误不排队、失焦静默以及前台恢复后新错误正常显示。1440 像素桌面和 390 像素手机浏览器测试通过，并验证后台错误仍可复制、复制日志脱敏、手动关闭、剪贴板降级与文档入口。')
p('浏览器计时采用可控时钟；无界面浏览器会模拟标签页焦点，因此后台测试显式模拟页面可见性和失焦事件，验证真实 DOM 事件处理与显示条件，不把该测试称为真人切换窗口验收。生成物一致性及修改文件 ESLint 检查通过。')
h('发布与使用')
p('修改前核对本机与远程 Main 提交一致。提交仅包含本次修复、版本、文档与测试，保留其他任务的工作区改动。Main 推送触发网站生产部署；刷新网页后使用新规则。在线说明位于 https://jackyun.top/docs/errors.html。')
for element in [doc.element,*[style.element for style in doc.styles]]:
    for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='JackYun 错误提示显示规则修复说明';doc.core_properties.author='Codex'
out=Path(__file__).resolve().parent/'JackYun-3.22.1-Error-Notice-Lifecycle-20261014.docx';doc.save(out);print(out)
