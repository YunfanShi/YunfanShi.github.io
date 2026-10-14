from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).resolve().parent;repo=out.parents[1];app=repo/'jackyun-portal'
script=(app/'public/shared/error-notices.js').read_text(encoding='utf-8')
catalog=json.loads(re.search(r'window.JackyunErrorCatalog=(.*);\n',script).group(1))
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name,size in [('Normal',11),('Title',23),('Heading 1',16),('Heading 2',12)]:
    style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('JackYun 错误系统   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,2)
def page(title):
    heading=doc.add_heading(title,1);heading.paragraph_format.page_break_before=True
doc.add_paragraph('JackYun 全站错误系统与同步部署说明',style='Title')
p('网站版本 3.20.0    同步后台版本 3.5.0    2026 年 10 月 14 日')
p('本次部署同步压缩升级，并将原来主要存在于日志中的错误转成用户可直接看到、可定位和可复制诊断的右上角提示。覆盖网站根页面、模块、管理员页面和独立 HTML 页面；25 个稳定错误码对应同一份在线文档。')
h('用户可见行为')
p('发生错误时，上方显示中文说明，下方显示 JY 错误码，错误码下面提供复制日志按钮和错误码说明链接。提示采用右上角布局，按照本次需求后半部分的具体说明执行。桌面宽度最多 380 像素，手机按屏幕宽度收缩；不抢走正在输入的焦点，用户可以关闭提示。')
p('同一错误码和来源在十秒内出现多次时合并计数，最多同时显示三条。系统不因重试不停新增同一提示；关闭弹窗不会清除日志。主动取消 fetch 或 XMLHttpRequest 不弹错误提示，超时和真实网络失败分别使用对应编号。')
h('同步兼容与自动转换')
p('服务器同时接受原始 JSON 分片与 Gzip 分片。新客户端读取服务器的能力声明后自动先压缩再上传；下载时请求压缩分片，逐片解压后校验完整内容。旧客户端继续使用原格式，用户无需导出、重新导入或手动判断传输格式。后续新上传的临时分片自动采用压缩格式。')
p('业务表继续保存可查询的 JSON、版本号、内容哈希和删除标记；没有批量替换业务内容，也没有数据库迁移。自动转换针对传输格式，不能把旧设备的快照当作最新内容。写入仍需匹配服务器版本和哈希，双方独立变化自动合并，同字段真实分歧保留副本。')
h('线上部署证据')
p('Azure 应用 jackyun-sync-9d26b2cc 的最终发布编号为 2d66de66-ff9c-4345-8082-f3e714ef9873。CLI 返回 Deployment was successful；公开健康接口返回 version 3.5.0、syncProtocol 3、transferEncoding gzip、legacyEncoding identity。此证据证明后台版本和格式兼容已生效，不代表已用真实账号完成两台设备的生产数据验收。')
p('网站通过 Main 的 Git 集成发布到 Vercel 的 jackyun-portal 项目。最终线上发布编号及浏览器检查记录由发布时核对；本说明与代码一同进入 Main，保留其他任务的未提交改动。')

page('技术实现与验证范围')
h('全站捕获入口')
p('共享运行脚本在网站根布局加载，覆盖登录页面及所有子路由；旧 HTML 通过既有 site dialog 脚本加载，未使用该脚本的页面直接加载错误运行脚本。PWA 缓存升级到 v7，预缓存错误脚本和错误码文档，离线页面仍可显示错误说明。根布局严重错误另有 global error 页面作为最后边界。')
p('捕获 fetch 和 XMLHttpRequest 的 HTTP 错误、网络异常、console error、未捕获脚本错误、Promise 拒绝、资源加载失败以及模块的诊断事件。HTTP 200 但 ok false 或 success false 的标准 JSON 响应也会识别。额外读取的 JSON 响应最多 16000 字节，避免复制解析巨大的成功同步正文。')
p('业务代码需要主动发出诊断或返回标准错误响应；被完全吞掉、既无响应也无日志的错误无法凭空识别。Next 的服务端错误通过页面边界或 HTTP 响应送达浏览器。对未登记的后台字符串编号保留原编号，同时按 HTTP 状态、功能标签和异常类型映射到文档化的 JY 编号。')
h('日志和向后兼容')
p('网站 API error 保留原始 code、message 和 requestId，并新增 siteCode。Azure 错误响应保留原来的字符串 error 和 code，也新增 siteCode。旧调用者不需要改变读取方式；新提示使用统一编号，原始技术原因和请求编号仍保留在复制的日志中。')
p('复制内容包含当前错误、时间、路径、重复次数、原始 API 编号、请求编号、HTTP 状态和近期诊断。复制前再次脱敏口令、令牌、Cookie、API Key、凭据和会话字段，并去掉 URL 查询参数。不会自动上传这些日志。浏览器禁止剪贴板时，弹窗提供只读文本框供选择和手动复制。')
h('验证与实际限制')
p('125 项同步回归和 17 项后台回归通过；新增错误系统三项测试验证编号映射、文档完整性、浏览器生成物一致性、重复错误上限、主动取消及 31 个独立 HTML 页的接入。桌面 1440 像素和手机 390 像素浏览器验证通过，包括复制、脱敏、手动复制、关闭、输入焦点和在线文档。')
p('本机生产页面实测：根登录页面接收真实 API 错误后弹出对应编号；PracticeTimer 独立页面的模块诊断也显示错误。由于本机尚未配置同步转发地址，本机同步 API 返回 503 和 JY-2101；它验证错误链路，不被当作线上同步故障。类型检查、修改文件 ESLint、生产构建和私钥泄漏检查通过。')

def code_table(entries):
    table=doc.add_table(rows=1,cols=3);table.autofit=False
    widths=[1.05,3.75,2.2]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,text in zip(table.rows[0].cells,['错误码','显示说明','建议处理']):cell.text=text
    rowprops=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');rowprops.append(repeat)
    for code,entry in entries:
        row=table.add_row()
        for cell,text in zip(row.cells,[code,entry['message'],entry['action']]):cell.text=text
    for row in table.rows:
        for cell,width in zip(row.cells,widths):cell.width=Inches(width)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(6);para.paragraph_format.space_before=Pt(4)
                for run in para.runs:run.font.size=Pt(10.5)
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:run.bold=True
page('错误码目录一')
p('账号、网络、服务及同步错误。具体后台别名保存在在线开发文档中；复制日志可以查看原始编号和请求编号。')
items=list(catalog.items());code_table(items[:13])
h('维护方法')
p('目录源文件为 src/lib/error-codes.ts。修改后运行 node scripts/build-error-system.mjs；--check 检查生成物。npm run verify 已加入目录与编号测试。已发布编号保持语义稳定，新增问题分配新编号。')
page('错误码目录二')
p('输入、存储、页面、AI、集成和兜底错误。未知技术错误仍有可反馈编号，不会只剩下控制台文字。')
code_table(items[13:])
h('文档与发布检查入口')
p('在线错误码说明 https://jackyun.top/docs/errors.html。开发文档位于 jackyun portal docs error codes md，浏览器可下载 /docs/error-codes.md。同步后台健康接口为 https://jackyun-sync-9d26b2cc.azurewebsites.net/health/live，仅提供运行元数据。')
p('Vercel 发布时核对 team kcV9DOM11BzvUqrfp1TviqOU、项目 jackyun portal、Main 提交及 production READY 状态。后续回滚应回滚程序部署，保留业务数据库版本；新版后台继续接受旧格式输入。')
for element in [doc.element,*[style.element for style in doc.styles]]:
    for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='JackYun 全站错误系统与同步部署说明';doc.core_properties.author='Codex'
target=out/'JackYun-3.20.0-Error-System-Deployment-20261014.docx';doc.save(target);print(target)
