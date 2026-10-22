from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out=Path(__file__).parent
doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(2)
sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
for st in doc.styles:
 for n in list(st.element.iter()):
  if n.tag in [qn('w:pBdr'),qn('w:u')]:n.getparent().remove(n)
  elif n.tag==qn('w:color'):
   n.set(qn('w:val'),'000000')
   for a in ['themeColor','themeTint','themeShade']:n.attrib.pop(qn('w:'+a),None)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(14)
def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,1)
doc.add_paragraph('雅思全屏修改与外部 AI 工具优化说明',style='Title')
p('JackYun 3.27.0    2026 年 10 月 19 日')
p('本次更新将屏幕空间优先交给作文编辑和批改工具：移除顶部修改完成了吗整块提示，将保存最终稿放入编辑工具栏，新增全屏修改，并恢复无需触发站内 AI 报错即可使用的外部 AI 导入。')
h('常规页面的调整')
p('工作台顶部不再显示完成提示及说明卡片。保存最终稿按钮与全屏修改按钮放在作文编辑区的布局工具旁边，保留字数和保存状态。普通左右分栏增加可用高度，最终稿文库与阅读器的入口保留。')
h('全屏修改的范围')
p('点击全屏修改后，作文编辑和修复工具占满屏幕。桌面左右布局让正文与反馈并排滚动，手机采用纵向滚动。页面大标题、导航和其他站点区域不进入全屏内容。')
p('全屏内保留正文输入、题目折叠、字数、保存状态、布局切换、高亮、AI 分析、语言升级、外部提示词和回复导入，以及批改反馈。修订地图、迁移说明和草稿安全卡片暂时收起，退出后恢复。进入全屏时滚动区域回到顶部。')
p('题目在全屏中默认折叠，需要查看或修改时点击题目展开。题型选择在普通编辑视图中调整。顶部小工具条保留计时开始或暂停、剩余时间、保存最终稿与退出全屏，不再放置大块使用说明。')
p('优先使用浏览器原生全屏。如果浏览器不支持或拒绝全屏，则使用占满浏览器视口的专注工作区。退出按钮两种模式均可用，原生模式遵循浏览器 Esc 行为，窗口内模式也支持 Esc。退出恢复页面滚动。')
h('保存和反馈优先级')
p('从全屏点击保存最终稿时，仍要求至少 80 个字符。保存成功后退出全屏并进入阅读器；文字不足时继续留在编辑区并显示提示。全屏切换不创建新文章，不重置草稿、计时或反馈。')
p('外部 AI 面板始终有清晰入口。没有反馈时默认展开，方便复制提示词并导入；已有批改反馈时自动收起输入区，给问题和修改建议让出空间，需要再次导入时可重新展开。站内调用失败的备用状态仍展开并显示相应说明。')
doc.add_page_break()
h('编辑后的引用跟随')
p('高亮不再只用旧句做精确匹配。每次收到站内或外部反馈，记录对应作文与批改时的文本；之后把引用范围映射到当前修改稿。删除、插入或替换部分词语时，剩余对应内容继续高亮，反馈卡片与语言升级卡片的当前引用也同步更新。点击定位选中当前修改后的范围，撤销文字或返回上一轮反馈时恢复对应锚点。')
p('例如将 the advantages for taking risks in the personal lives and professional lives can overwhelm the drawbacks 改为 the advantages of taking risks in personal and professional lives outweigh the drawbacks，高亮和引用跟随新句。完全删除或失去可靠对应时显示无法定位提示，不把其他无关句子误标。原始 AI 反馈与锁定原稿保留；引用跟随不表示自动判定已经改对，仍可标记已改或进行复查。')
h('外部 AI 导入的操作')
p('进入写作工作台，在修复循环中展开外部 LLM。填写题目与作文后复制完整提示词，在外部 AI 取得符合 JSON 规范的回复，再粘贴并点击导入 AI 回复。不需要配置站内 AI，也不需要先故意触发调用错误。')
p('旧代码将整个导入面板限制为 externalFallbackMode 存在时显示。本次移除这个显示限制，保留备用场景文案；为回复输入框增加可访问名称，站内 AI 请求期间及草稿初始化时禁止导入，避免结果相互覆盖。新建文章仍清空旧的临时反馈。')
h('实现与版本')
p('主要修改 writing-workbench.tsx 的工作区结构、全屏状态与事件清理，继续使用既有草稿保存和文库逻辑。增加 fullscreenchange 与 Escape 监听，卸载时移除；专注模式结束后恢复原页面滚动设置。没有新增数据库迁移、外部依赖或模型请求。')
p('开始任务时，本地 main 与 origin/main 均为 912f56b，版本 3.26.2。本轮外部 AI 入口修复和随后提出的全屏优化合并为一次交付；因增加全屏功能，按新功能规则更新到 3.27.0，并同步应用版本、锁文件与站内更新记录。')
doc.add_page_break()
h('验证范围和结果')
p('28 项写作单元及并发保存回归通过，目标文件 ESLint、TypeScript 和生产构建用于检查兼容性。浏览器使用真实工作台与本地存储逻辑，覆盖 1440 像素英文与 390 像素中文视口。')
p('浏览器验证原生全屏占满屏幕、编辑器和工具可用、辅助卡片隐藏、计时按钮、退出按钮、原生全屏失败后的视口专注及 Escape，以及从全屏保存后进入阅读器并恢复页面滚动。验证不触发站内失败即可查看提示词并成功导入 JSON，随后新建文章会清空旧反馈。')
p('原有新建、快照恢复、保存失败保护、阅读器分段、三个阅读主题、长标题及正文保存回归继续执行。结构化结果和桌面手机截图保存在 verification 目录。测试使用本地夹具，外部 AI 和导航替换为测试实现，未将结果解释为真实账号云端同步或线上部署完成。')
p('新增引用跟随回归覆盖用户提供的长句删词与替换、前置插入、重复引用、邻接插入、整段删除和无关替换。浏览器检查编辑器高亮、反馈引用、点击定位与文字撤销同步，避免引用在局部修改后直接消失。')
h('数据与交付')
p('已有草稿、文库和错误历史继续使用原存储键。外部反馈保持原有临时状态语义，刷新后不作为独立报告自动恢复。代码、版本、测试、验证材料和本说明直接提交并推送 main；需要回退时撤销此次提交并按既有流程部署。')
doc.core_properties.title='雅思全屏修改与外部 AI 工具优化说明';doc.core_properties.author='JackYun'
file=out/'雅思全屏修改与外部AI工具优化说明3.27.0.docx';doc.save(file);print(file)
