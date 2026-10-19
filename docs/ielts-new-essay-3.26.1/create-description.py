from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out=Path(__file__).parent
doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(2)
sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
for st in doc.styles:
 for n in list(st.element.iter()):
  if n.tag in [qn('w:pBdr'),qn('w:u')]:n.getparent().remove(n)
  elif n.tag==qn('w:color'):
   n.set(qn('w:val'),'000000')
   for a in ['themeColor','themeTint','themeShade']:n.attrib.pop(qn('w:'+a),None)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(14)
def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,1)
doc.add_paragraph('雅思写作新建文章修复说明',style='Title')
p('JackYun 3.26.1    2026 年 10 月 19 日')
p('本次修复解决雅思写作修复工作台难以开始下一篇文章的问题。在页面顶部增加新建文章按钮，从写作工作台、最终稿文库和阅读器都能直接进入空白编辑器。切换前保存上一篇草稿的恢复快照，避免为了写新文章而手动删除原文。')
h('问题原因与影响')
p('旧页面仅在工作台下方迁移验证卡片中提供开始新题迁移入口，顶部没有通用的新建文章按钮，文库和阅读器也没有直接开始下一篇的操作。旧的切换逻辑把新空白稿写入快照，没有专门保存离开的旧稿，而且没有清除外部 AI 回复与备用模式。')
p('此次将普通新建与迁移练习统一到 startNewEssay 流程。新建不要求先完成 AI 批改，也不要求达到保存最终稿所需的 80 字符门槛；只有题目或一小段正文的草稿，同样先保存恢复快照。')
h('使用方法')
p('进入 /ielts-writing，点击顶部新建文章，英文显示 New essay。页面自动返回写作工作台并聚焦正文输入框，可以直接填写题目和作文。文库与阅读器共用同一个顶部按钮，不必先手动返回工作台。')
p('如需找回刚刚离开的草稿，在草稿安全卡片点击恢复快照。快照沿用现有最多五份的轮换规则，并非无限篇文章文库。希望长期保留的完成稿，仍应先保存到最终稿文库。文库原有最终稿不会因新建文章被删除。')
h('新文章的状态')
p('每次新建生成独立 essayId，题目、正文和锁定原稿清空。保留当前 IELTS 题型，Task 1 重置为 20 分钟，Task 2 重置为 40 分钟，计时器保持暂停。批改反馈、反馈撤销记录、当前问题、高亮选择、已改标记、上轮规则及外部 AI 回复状态归零；错误历史及最终稿文库保留。')
doc.add_page_break()
h('保存与并发保护')
p('切换顺序为先保存旧稿和恢复快照，再持久化新稿，最后更新页面状态。使用现有三方合并写入逻辑，继续保留其他标签页的独立修改，并由现有冲突检查拒绝冲突覆盖。')
p('为本次切换增加严格失败模式。主稿、镜像、恢复快照或冲突检查出现异常时终止切换，显示保存错误并保留当前编辑内容；不把保存失败伪装成新建成功。普通输入保存继续使用现有行为。')
p('页面初始恢复完成之前，以及站内 AI 分析期间，新建文章与开始新题迁移按钮均禁用。处理函数也检查这两个条件，防止未恢复草稿即切换，或旧文章的分析结果写入新文章。')
h('改动范围与版本')
p('修改 writing-workbench.tsx 的入口、状态重置和严格保存分支。同步更新 package.json、package-lock.json、APP_VERSION 和站内更新记录，版本从最新主线 3.26.0 升到 3.26.1，按小问题修复规则增加最后一位。')
p('任务开始时本地 HEAD 为 acaa314，在线核对远端 main 为 4a4c042，版本为 3.26.0。本次先备份修复内容，将本地 main 快进到远端，再在最新代码上重新实施修复。没有创建功能分支。')
h('验证方法')
p('写作回归包含 24 项单元测试，覆盖反馈解析、原稿比较、镜像恢复、文库及实际组件保存处理中的跨标签页合并。目标代码 ESLint、私钥标记扫描和生产构建通过；生产构建包含 TypeScript 检查。')
p('新增可重复执行的浏览器回归 tests/ielts-new-essay.browser.mjs，直接挂载真实 WritingWorkbench 组件及本地持久化逻辑，使用 1440 像素英文与 390 像素中文视口。检查顶部新建、独立编号、空白字段、题型计时、旧稿快照、刷新恢复、恢复旧稿、文库与阅读器入口、迁移练习、保存失败保留和 AI 处理中禁用。结果与截图保存在 verification 目录。')
p('浏览器夹具将外部 AI 请求和导航链接替换为测试实现，没有调用真实 AI 或使用用户云端账号。本次验证确认本地组件行为，线上部署完成与真实账号跨设备同步需以部署平台及实际访问结果为准。')
h('交付与回退')
p('源码、浏览器回归脚本、验证结果和本说明一并提交 main 并推送远端。若需回退，撤销本次修复提交并按既有发布流程部署上一版本。已有草稿、恢复快照和最终稿继续沿用原有存储键，不需要数据库迁移。')
doc.core_properties.title='雅思写作新建文章修复说明';doc.core_properties.author='JackYun'
file=out/'雅思写作新建文章修复说明3.26.1.docx';doc.save(file);print(file)
