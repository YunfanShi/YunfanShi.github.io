from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out=Path(__file__).parent
doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(2)
sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
for st in doc.styles:
 for n in list(st.element.iter()):
  if n.tag in [qn('w:pBdr'),qn('w:u')]:n.getparent().remove(n)
  elif n.tag==qn('w:color'):
   n.set(qn('w:val'),'000000')
   for a in ['themeColor','themeTint','themeShade']:n.attrib.pop(qn('w:'+a),None)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(14)
def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,1)
doc.add_paragraph('雅思阅读器段落显示修复说明',style='Title')
p('JackYun 3.26.2    2026 年 10 月 19 日')
p('本次修复针对雅思写作最终稿阅读器的正文分段不明显、深色网站主题下标题输入框文字难以辨认，以及长标题占据过多页面空间的问题。阅读器现在按原文换行显示独立段落，并提供明确段间距；文章原文和文库记录继续保留。')
h('问题定位')
p('旧阅读器使用单个 div 配合 whitespace-pre-wrap 显示全文。它能够保留部分换行，但没有独立段落结构与段间距，连续输入或粘贴的段落容易看成一整块正文。原稿即使只使用单换行分段，也需要在阅读模式中显示清楚。')
p('截图中的标题编辑框使用网站背景色，文字却继承了阅读器的配色。深色网站配合纸色阅读器时，形成深底深字。本次明确指定输入文字使用与背景相配的站点前景色。')
h('分段显示规则')
p('新增 splitWritingParagraphs 纯函数，兼容 Windows 的 CRLF、LF、旧式 CR，以及 Unicode 行分隔符和段分隔符。每个非空文本段以独立 p 元素显示，段间距为当前字号的 1.25 倍，字号和行距调整继续生效。空白行不会产生额外空段。')
p('渲染时去除每段两侧空白，段内连续空格保留。显示逻辑不改写归档内容，不插入或删除文章词句，不重新计算或覆盖存储中的正文。刷新或重新打开旧文章即可使用新的显示方式，不需要重新保存最终稿。')
p('如果保存的原文完全没有换行，仍显示为一个段落。本次不会根据句号或语义猜测段落边界，以免替用户改变文章结构；这类文章应在原稿中加入所需换行后重新保存最终稿。')
h('标题与主题显示')
p('标题使用适合阅读的字号：手机为 24 像素，较大视口为 30 像素，并允许长词或长链接换行。完整标题保留，不通过截断隐藏文字。标题编辑框显式使用站点前景色，在纸色、薄荷色及夜间阅读主题中均保持可辨认。')
doc.add_page_break()
h('实现与改动范围')
p('ielts-writing.ts 增加段落拆分函数；writing-workbench.tsx 在阅读器正文中生成独立段落并调整长标题布局；archive-title-editor.tsx 修复输入框文字颜色。没有增加 AI 调用、网络请求或数据库迁移。标题保存继续使用既有文库更新与合并逻辑。')
p('本次按 React 组件检查要求复核语义结构、列表键、状态与渲染行为。段落使用 p 标签；列表仅按不可编辑的正文顺序渲染，索引键不会造成输入状态错位；没有增加副作用或事件监听器。')
h('验证结果与范围')
p('写作单元及并发保存回归共 25 项通过。新增段落测试覆盖 CRLF、单换行、多个空行、CR、Unicode 分隔符、空白正文、单段多句与段内连续空格。目标文件 ESLint 和生产构建检查用于确认类型与打包兼容性。')
p('浏览器回归直接挂载真实 WritingWorkbench、ArchiveTitleEditor 和本地存储逻辑。使用 1440 像素英文与 390 像素中文视口，通过文库的打开阅读器按钮进入实际阅读视图，验证四段内容和明显段间距。')
p('在模拟深色网站配色下逐一切换纸色、薄荷色及夜间阅读主题，检查输入框前景和背景颜色、段落数量与页面宽度。编辑并保存长标题后，检查完整标题可见、页面无横向溢出，并确认保存的正文内容与测试原文完全一致。')
p('原有新建文章回归同步执行，继续检查草稿快照、刷新、恢复、新建、迁移练习和保存失败保护。浏览器使用本地组件夹具，外部 AI 与导航替换为测试实现；未使用真实云端账号，也不将本地结果作为线上部署完成的证明。')
h('版本与仓库交付')
p('任务开始时，本地 main 与获取后的 origin/main 均为 875d5de，版本为 3.26.1。本次按小问题修复规则升级为 3.26.2，更新 package.json、锁文件、APP_VERSION 与站内更新记录，直接在 main 提交并推送源码、回归测试、验证材料和本说明。')
p('已有文章继续使用原来的存储键。若需回退，可撤销本次修复提交并重新部署上一版本；本次没有对用户已有正文做迁移或覆盖。')
doc.core_properties.title='雅思阅读器段落显示修复说明';doc.core_properties.author='JackYun'
file=out/'雅思阅读器段落显示修复说明3.26.2.docx';doc.save(file);print(file)
