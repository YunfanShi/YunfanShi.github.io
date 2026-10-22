from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
root=Path('../docs/offline-recovery-3.24.5');d=Document();sec=d.sections[0];sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Microsoft YaHei';st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.22
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(15)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
d.add_heading('JackYun 数据库恢复与离线使用修复说明',0)
p('版本 3.24.5    日期 2026 年 10 月 17 日    范围 网站与 PWA')
p('本次修复解除 PRACTICE TIMER 对首次云端连接的依赖。已保存到设备的页面在断网或服务器错误时继续打开；练习、逐题计时、本机保存、刷新恢复和报告生成经过桌面与手机尺寸的浏览器验证。页面不再反复显示连接失败或离线模式提示，真实保存错误和设备接管冲突仍可见。')
h('数据库故障与恢复结果')
p('排查时，项目管理接口显示健康，但实际 SQL 查询返回 Connection terminated due to connection timeout。访问日志存在认证令牌刷新 504，数据库日志出现 statement timeout。这些证据说明项目状态标签不能替代实际连通性检查；尚无充分证据断定故障由某条业务查询或数据损坏引起。')
p('用户明确批准生产项目重启及数分钟不可用的影响后，控制台显示重启进行中。重启后重新执行 SQL，业务表 profiles 与 practice_timer_live 均可读取，数据库不处于恢复状态，闲置事务数量为 0。认证健康接口与 profiles 的限零行 REST 查询均返回 HTTP 200。整个恢复过程未清空、重建或删除业务数据，也未修改数据库权限和策略。')
h('流量配额的独立问题')
p('用户提供的用量页显示：本计费周期为 2026 年 9 月 25 日至 10 月 25 日，Egress 为 6.08 GB，免费额度为 5 GB，超出 1.08 GB。控制台提示组织处于宽限期，若仍然超额，2026 年 11 月 8 日起可能受限，受限请求会返回 402。该提示不能解释本次已观测到的 SQL 超时，重启也不会重置用量。未执行付费升级或变更套餐。')
d.add_page_break()
h('计时与本机保存的修改')
p('原行为：首次云端读取失败后，计时控制器进入 unavailable 状态；新建或恢复练习会再次等待连接，失败便拒绝开始。现行为：云端发现可在后台执行，loading、unavailable、未登录或浏览器确认断网时均允许本机练习；已经连接时发生临时 claim 错误也可回落到本机。服务器明确返回的冲突仍按原流程处理。')
p('计时中断网不会停止本机时钟。已有控制令牌的练习在重连后继续尝试保存最新计时，再释放控制权；只有服务器明确确认丢失控制权才退出。无令牌的本机练习不会在重连时自动接管其他设备或覆盖已有云端练习；若云端没有所有者和任何会话负载，可取得控制并保存。普通 Schedule 数据继续沿用原有本机同步和冲突处理流程。')
p('原有断网恢复副本已能保存，但刷新后没有读取入口。本次在计时器桥接层读取并校验恢复副本，仅在其时间不早于常规记录时纳入恢复列表。正常持久化成功后才清理对应恢复副本。无效副本保留以便人工恢复；不会把未保存成功的内容显示成已经同步云端。')
h('PWA 页面与资源缓存')
p('Service Worker 缓存版本升级为 v8。补齐 /practice-timer 门户页面与 /PracticeTimer.html iframe 导航，安装阶段预存计时引擎、音频逻辑、云端控制脚本、工作台脚本以及三份样式文件。导航遇到网络失败或服务器 5xx 时，优先返回相同地址的有效缓存。API 保持网络请求，不缓存认证或业务接口响应。')
p('门户的离线链接处理同步加入 /practice-timer，避免断网时依赖 Next.js 在线导航请求。账号变化清理个人页面、同地址与查询参数匹配、页面有效期和缓存数量限制继续保留。')
h('页面展示')
p('计时器隐藏常规云端连接状态，仅以“已保存”反馈本机保存确认。移除全局连接故障横幅，顶部保留中性的“待同步”入口，便于需要时检查详情。未缓存页面改为说明该页尚未保存在设备上，并提供返回已保存工作区的入口。明确的设备接管、实际保存失败和数据冲突仍保留提示。')
d.add_page_break()
h('验证内容与结果')
p('自动化测试：计时引擎、课表关联、预设时限、训练统计、PWA 清单、静态资源、导航缓存、断线控制器、计时桥接层和云端路由相关测试通过。新增覆盖首次断网开始、失败 claim 回落、无令牌重连安全性、刷新恢复检查点，以及服务器 503 返回已缓存计时页面。')
p('真实浏览器：在 Edge 无头浏览器的 1440 像素桌面与 390 像素手机视口运行正式构建。先联网缓存页面，再切断网络，新建练习、运行计时、保存、刷新恢复、暂停继续、结束查看报告及恢复联网均通过，未出现页面脚本异常。截图和浏览器结果存放在同目录 verification。')
p('跨设备模拟：通过本地服务模拟云端控制接口，验证每分钟检查点、下一题开始、暂停冻结、设备接管、断线继续、重连补传与旧设备令牌失效。此项验证的是前后端控制契约，不等于两台真实设备的生产账号实测。')
p('构建与代码检查：Next.js 16.3.8 正式 webpack 构建通过，构建包含 TypeScript 检查；改动的 React 组件 ESLint 检查通过。复核了 hooks 清理、事件监听释放、异步竞态、消息来源校验和异常数据保护。')
h('发布与使用条件')
p('本次属于现有行为修复，版本由 3.24.4 更新为 3.24.5，package.json、锁文件与站内更新记录保持一致。修改基于远端 main 最新提交的隔离检出完成，不创建功能分支，也不混入主工作目录的其他未提交内容。验证后将本次修复与本文档同步到 main。')
p('已有安装需要先联网载入新版本并完成 Service Worker 更新。已访问并保存的本地工作页面可离线使用；浏览器清理存储、缓存被回收或从未访问的页面无法凭空离线生成。现有个人页面缓存期限为 7 天、最多 24 个地址。AI 生成、第三方调用、首次下载和账号管理仍需要对应服务。')
p('本次验证确认数据库恢复和离线修复有效；生产部署完成后仍应核对发布提交与用户设备上的 Service Worker 版本。重启只恢复本次服务可用性，尚不能证明潜在的资源瓶颈已长期消除。')
for st in d.styles:
 for node in list(st.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 if st.name.startswith(('Title','Heading','heading')):
  st.font.color.rgb=RGBColor(0,0,0)
for para in d.paragraphs:
 for run in para.runs:
  run.font.color.rgb=RGBColor(0,0,0)
d.save(root/'数据库恢复与离线修复说明3.24.5.docx')
print(root/'数据库恢复与离线修复说明3.24.5.docx')
