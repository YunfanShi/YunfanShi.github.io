from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
root=Path(__file__).parent
d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1']:
 st=d.styles[name];st.font.name='Microsoft YaHei';st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(8);d.styles['Normal'].paragraph_format.line_spacing=1.2
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
d.add_heading('JackYun 计时冲突修复说明',0)
p('网站 3.25.1 · 2026 年 10 月 17 日 · Azure 后台保持 3.6.0')
h('问题与证据')
p('用户提供的诊断包含六次 /api/practice-timer/live HTTP 409，但未记录 conflict 字段；因此不能仅凭附件判定每次冲突的来源。只读聚合显示当前有一份未完成云端草稿，最后心跳租约已过期数小时。没有读取练习正文，也没有删除或修改生产草稿。')
p('源码可确认两项缺陷：不同练习编号遇到已有未完成草稿时，服务器返回 ACTIVE_SESSION，客户端直接抛出并阻止开始；通用错误收集器又将无明确错误码的 409 分类为输入错误 JY-3001。旧冲突响应可能附带很大的草稿，使有限大小的诊断解析也无法取得原因。')
h('行为调整')
p('遇到 ACTIVE_SESSION 后，新练习继续在本机开始。清除本次云端控制令牌和延迟释放标记，保留本机练习编号与待保存状态；不向已有草稿发送 save 或 release。原有云端草稿继续保留，并通过“接续云端练习”入口打开。没有自动覆盖、完成或删除旧草稿。')
p('本机练习经原 Schedule 校验和合并流程保存，沿用 Azure 工作区同步；离线时使用原本机保存与恢复机制。该练习不会强行占用云端实时计时槽，因此不提供同一份实时计时的跨设备接续，直到用户主动接续原云端练习或实时槽空闲。')
p('API 对已知冲突返回精简的 conflict、JY-2103、具体 message 和 handled 标记，不再返回草稿正文。错误收集器仅对该计时端点、HTTP 409、三个明确已处理原因取消通用弹窗，仍记入脱敏诊断。其他接口的 409 和未知冲突继续提示。LOST_CONTROL 仍由计时器展示设备接管并停止旧设备写入；STALE_REVISION 仍按原流程重新读取并重试。')
d.add_page_break()
h('验证与边界')
p('计时相关回归 39 项通过；错误码与生成物检查、密钥扫描、改动 TypeScript 文件 ESLint、正式生产构建和 TypeScript 检查通过。浏览器验证包括：桌面计时与暂停、分钟检查点、跨设备接续、断网继续、重连上传，以及手机视口在已有草稿时创建独立练习并保存，旧草稿与服务端保存次数保持不变。')
p('另在桌面和手机视口验证：已处理计时冲突不弹通用错误；复制其他错误的日志可看到 ACTIVE_SESSION；其他端点 409 和计时未知冲突仍弹窗。上述浏览器为受控接口测试，不等同于两个真实生产账号的端到端验收。生产发布后仅进行无需登录的资源与版本检查，不伪造用户令牌。')
p('PWA 缓存版本从 v8 升至 v9，让联网后的更新可以预缓存新计时脚本。已经完全离线且从未获得更新的设备仍使用原缓存，需要联网下载一次新版本。')
h('登录有效期的建议')
p('保持登录会话与访问令牌有效期是两件事。本次不修改 Supabase JWT 有效期。建议长期保留登录，并自动刷新短期访问令牌；正常请求通过已缓存公钥本地验签，避免每次向 Auth 请求用户资料。刷新或云端不可用时，本机练习继续，恢复连接后再同步。')
p('24 小时访问令牌能够减少刷新频率和刷新服务故障的影响，但被盗令牌也可能在剩余有效期内继续使用，登出并不会让已签发令牌立刻失效。学习网站仍有个人记录与文件，因此不把“学习用途”视为可以取消签名、有效期或行权限校验的理由。Supabase 官方建议通常不超过一小时，具体可按可用性和撤销需求取舍。')
h('发布与回退')
p('本次为错误修复，按规则从 3.25.0 升至 3.25.1。复用原工作目录，不建立新分支；先验证再快进合并 main，并保留主目录已有未提交修改。通过 main 原发布流水线部署网站。本次没有新增数据库迁移、付费资源或 Azure 部署。回退可部署前一网站版本，生产草稿和原数据库函数均不受影响。')
p('前一发布 3.25.0 / Azure 3.6.0 已完成错误关闭后同类静默 30 分钟、本地验签、Azure 账号核验缓存与计时元数据减负。主工作区同步已在 Azure；剩余文件、账号状态、通知与业务数据仍需逐模块追平、校验和切换，不能直接使用历史快照替换当前权威数据。详细方案见仓库 backend-load-and-errors-3.25.0 文档。')
p('参考：https://supabase.com/docs/guides/auth/sessions ；https://supabase.com/docs/reference/javascript/auth-getclaims')
for st in d.styles:
 for node in list(st.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 if st.name.startswith(('Title','Heading','heading')):st.font.color.rgb=RGBColor(0,0,0)
for para in d.paragraphs:
 for run in para.runs:run.font.color.rgb=RGBColor(0,0,0)
d.save(root/'计时冲突修复说明3.25.1.docx')
