from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).parent
doc=Document()
sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(2.0)
sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for st in doc.styles:
 for node in list(st.element.iter()):
  if node.tag in [qn('w:pBdr'),qn('w:u')]:
   node.getparent().remove(node)
  if node.tag==qn('w:color'):
   node.set(qn('w:val'),'000000')
   for attr in ['themeColor','themeTint','themeShade']:
    node.attrib.pop(qn('w:'+attr),None)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.25
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,1)
doc.add_paragraph('练习计时器与雅思四科计时更新说明',style='Title')
p('JackYun 3.26.0    2026 年 10 月 17 日')
p('本次更新完成三个改进：在学习侧边栏直接打开练习计时器；答题途中可以更正题型并保留已有时间；增加雅思听力、阅读、写作、口语四个板块的分段计时。雅思练习直接选择板块即可使用，不需要先在 Schedule 课表中建立雅思科目。')
h('入口与开始方式')
p('学习侧边栏新增练习计时器，英文界面显示 Practice Timer，直接打开 /practice-timer。进入后侧边栏能够识别当前页面。原有日程及时间管理入口继续可用，手机通过学习导航抽屉进入相同页面。')
p('开始页新增计时内容选择。课表逐题练习沿用科目关联、普通、宽松和模考档位；选择 IELTS 雅思四科后显示板块选择及该板块的时间说明，隐藏无关的课表科目与节奏档位。仍需选择练习内容类型并填写名称，方便在历史记录中找到对应材料。')
h('答题中更正题型')
p('普通逐题练习在开始计时后仍可通过顶部题型按钮、专注模式题型菜单或数字快捷键更正当前题型。切换时先结算当前已用时间，再更新题型和限时。已有用时、题目编号、名称、备注、批改结果和首次作答证据保留，总用时不归零。')
p('例如已经用了 3 分钟，原来误选 6 分钟题型，更正为 12 分钟题型后还剩 9 分钟；更正为 2 分钟题型则显示超时 1 分钟。运行中的题继续运行，暂停中的题保持暂停。题型选择框不暂停当前题；打开备注或设置面板仍按原规则暂停。')
p('下一题继续沿用先暂停当前题、选择下一题题型、选择后立即开始的流程。修改题型预设只影响后续题目，直接更正当前题型才会重新计算当前题的限时。已结束报告不提供继续作答或更换题型，首次批改快照不会被后续更正覆盖。')
doc.add_page_break()
h('雅思四科的计时规则')
p('写作按两个 Task 计时。Task 1 默认 20 分钟，Task 2 默认 40 分钟。可以从任意一个 Task 开始练习，点击下一个 Part 或 Task 切换；底部标签允许返回前一个部分，各部分分别保存已经使用的时间。')
p('阅读按三个 Part 计时，每个默认 20 分钟，总计 60 分钟。20 分钟是练习用的平均分配，不表示官方要求每篇固定限时。实际考试中阅读总时间固定，考生可根据文章和题目难度自行分配。')
p('口语按三个 Part 计时，默认 5、4、5 分钟。Part 1 和 Part 3 的官方时长各约 4 至 5 分钟；Part 2 的练习时间分为准备 1 分钟、陈述 2 分钟、追问 1 分钟，屏幕会按累计时间提示当前阶段。追问 1 分钟为练习安排，实际口语总时长约 11 至 14 分钟，由考官控制。')
p('听力的四个 Part 共用一个整场计时器，不把四段音频平均切成独立倒计时。默认 30 分钟作答加 2 分钟机考检查，总计 32 分钟。累计达到 30 分钟时显示最后 2 分钟检查阶段，达到 32 分钟后提示检查结束，并继续记录超时。')
p('不同听力录音的实际长度可能不同，本功能没有接入音频播放或自动识别音频结束。请按材料实际进度开始、暂停或结束；30 分钟是默认练习基准，检查阶段提示按已用时间触发，不代表录音一定在该时刻结束。机考的 2 分钟检查不使用纸笔考试的 10 分钟誊写时间。')
h('Part 切换与时间保留')
p('进入雅思练习时一次建立该板块的全部 Part 或 Task，所选部分成为当前部分。点击标签、下一部分按钮或 Part 选择菜单时，只切换到对应部分，不创建重复题目，也不覆盖已经使用的时间。若当前正在计时，切换后下一部分继续计时；若当前暂停，切换后保持暂停。')
p('听力只有一个整场计时单元，因此不显示下一部分按钮。阅读、写作及口语在最后一个部分隐藏下一部分按钮，可通过标签返回或结束查看报告。总时间是所有部分有效用时的累计。各部分超时不强制结束，也不会自动跳到下一部分。')
p('雅思使用固定配时，不受普通或宽松倍数影响；设置中保留声音、动效与颜色阈值。题目分值和每分分钟设置不用于雅思分段。报告仍可补充备注与批改，默认待批改不会被自动判断成答对或答错，也不自动计算雅思分数。')
doc.add_page_break()
h('保存和兼容处理')
p('新增可选雅思板块字段，区分听力、阅读、写作和口语，继续通过现有练习保存与读取流程处理；普通旧记录不需要迁移。保存验证器接受雅思配置，同时拒绝无效板块及与非雅思配置矛盾的数据。刷新恢复后默认暂停，保留已经保存的逐段时间。')
p('本次没有修改数据库结构及账号权限。登录后的同步、离线恢复和设备接管继续由既有计时同步模块负责。Service Worker 缓存版本从 v9 更新到 v10，重新预缓存 HTML 与计时脚本，避免新页面与旧脚本混用。离线页面仍需先在线访问和缓存。')
h('验证范围和结果')
p('单元回归共 68 项通过，覆盖计时结算、运行及暂停时更正题型、变短后的超时、首次作答证据保留、雅思各科配时与阶段边界、分段切换、保存验证、重新读取、导航高亮和离线缓存行为。原有课表绑定与练习统计相关测试同时执行。')
p('浏览器使用桌面 1440 像素和手机 390 像素视口检查。覆盖运行和暂停时改题型、专注快捷键、无课表时使用雅思四科、各 Part 的独立用时、刷新恢复、结束报告以及听力检查时间边界，检查无页面异常及横向溢出。浏览器保存请求经过真实的会话和配置验证函数，父页面与登录接口使用本地测试夹具。')
p('原有云端计时浏览器回归通过，覆盖下一题立即开始、暂停冻结、设备接管、离线继续、重连上传及已有云端草稿时开始本机练习。该回归模拟 HTTP 云端，没有使用用户真实登录账号做两台实体设备联调。本次不将本地测试结果解释为线上账号同步或真实音频自动计时已验证。')
p('生产构建、TypeScript 无输出检查及目标 TypeScript 文件 ESLint 检查通过，私钥标记扫描通过。功能截图和结构化浏览器结果保存在同目录 verification 中；浏览器自动化脚本与单元测试一同进入仓库。')
h('版本和仓库交付')
p('任务开始时本机 HEAD 与获取后的 origin/main 均为 3089386，版本为 3.25.1。本次新增雅思四科功能，按新功能规则将中间位加一，发布版本为 3.26.0，同步更新 package.json、锁文件和站内更新记录。在 main 直接提交本次源码、测试、详细说明及验证材料，保留其他任务未提交的工作区内容。')
h('配时参考')
p('IELTS Academic 考试结构：https://ielts.org/take-a-test/test-types/ielts-academic-test')
p('IELTS 口语结构：https://ielts.org/take-a-test/test-types/ielts-academic-test/ielts-academic-format-speaking')
p('IDP 机考检查时间：https://ielts.idp.com/uae/about/which-test-do-i-take/computer-delivered-ielts')
doc.core_properties.title='练习计时器与雅思四科计时更新说明'
doc.core_properties.author='JackYun'
doc.save(out/'练习计时器与雅思四科计时说明3.26.0.docx')
print(out/'练习计时器与雅思四科计时说明3.26.0.docx')
