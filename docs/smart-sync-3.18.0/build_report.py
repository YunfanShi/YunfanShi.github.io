from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

output = Path(__file__).resolve().parent
doc = Document()
section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.75)
for name,size in [('Normal',11),('Title',23),('Heading 1',16),('Heading 2',12)]:
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'; style.font.size=Pt(size)
    style.font.color.rgb=RGBColor.from_string('000000')
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('JackYun 同步系统   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def page(title):
    heading=doc.add_heading(title,1)
    heading.paragraph_format.page_break_before=True

doc.add_paragraph('JackYun 同步系统升级说明',style='Title')
p('应用版本 3.18.0    同步后台版本 3.4.0    2026 年 10 月 14 日')
p('本次针对多设备离线修改、旧数据先同步、冲突需要反复手动判断和大数据传输缓慢进行修复与升级。实现保留服务器版本校验和三方合并，新增压缩传输、冲突自动复核、持久操作编号和云端版本回退保护。说明同时记录实现边界、回归证据和上线顺序。')
h('问题与定位结果')
p('设备 A 的新增内容尚未上传，设备 B 持有旧快照并先同步时，服务器无法知道 A 存在未上传的内容。单纯按照谁最后同步、设备时间或设备自报的最大编号覆盖整页，都不能证明哪份内容完整。正确做法是保存每台设备最后确认的云端基准，仅上传本机实际变化，并在收到新版本后合并。')
p('当前代码已经具备每个存储键的服务器 revision、SHA256 内容哈希、按稳定 ID 比较的增量和三方合并。本次没有另造一套相互冲突的编号，而是补齐现有机制的缺口。原先的本机 sequence 在队列清空后会从小值重新开始；用户撤销修改并回到基准时，捕获逻辑可能忽略撤销，继续上传旧队列；已记录冲突也只刷新远端副本，不能自动发现冲突已经消失。')
h('版本与操作身份')
p('revision 是服务器在接受该存储键的新状态后推进的编号，用来判断客户端基准是否仍然有效。不同存储键有各自的版本，避免无关页面互相阻塞。baseRevision 与 baseHash 必须同时匹配，不能以设备时钟判断胜负。现有服务器对不匹配写入返回最新内容，由客户端重新合并后再提交。')
p('sequence 是本机产生操作的顺序，不用于比较两台设备谁的数据更新。它现在保存在 IndexedDB settings 中，与 outbox 写入同一事务；队列排空、页面重启和多个标签页同时编辑后继续递增。每次请求仍保留独立操作 UUID 和服务器回执，网络重试不会重复应用已接受的同一操作。')

page('多设备自动判断与内容保护')
h('同步执行顺序')
p('每次同步先捕获本机修改并持久化操作，再读取云端版本目录。版本和哈希一致的内容跳过正文下载。发现变化后读取该键的最新内容，将最后确认的基准、本机现值、云端现值进行三方比较。能够安全合并就保存备份、更新本机内容、重设提交基准，再上传增量；服务器确认后才移除队列中的对应意图。')
p('云端内容将要替换本机内容时，先保存本机备份并再次核对本机值。期间发生的新编辑不会被旧网络响应直接清掉。若下载的云端 revision 小于已保存的元数据版本或待上传基准，立即停止应用，保留本机内容和原基准，以免一次服务器回退变成永久覆盖。')
h('你提出的情景')
p('情景一：两台设备都从版本 1 开始。A 离线新增记录，B 未改内容先同步。B 不产生有意义的新修改，因此不会把 A 尚未上传的数据判为过期。A 重新联网后上传其新增记录，B 下次读取新版本即可得到相同内容。新增的回归测试验证云端没有因 B 的空同步产生一次写入。')
p('情景二：A 离线修改字段 a，B 修改字段 b 并先上传。A 同步时发现云端版本已前进，比较共同基准后自动保留 a 和 b 的变化，基于当前云端版本重新提交。双方按稳定 ID 新增不同记录也会合并；排序插入位置沿用现有合并规则。')
p('情景三：A 和 B 把同一个字段分别改成不同值。这时编号只能说明提交先后，不能说明哪份文字是用户真正想要的结果。系统保留双方副本及冲突位置，继续等待明确选择。删除与修改同一条记录、两侧不同重排、无法识别的普通数组变化仍按同样原则保护。')
h('减少手动判断的新增规则')
p('持久化冲突现在会使用当前本机内容和最新云端内容重新比较。若双方后来变得一致，或冲突字段恢复为可安全合并状态，自动继续同步并发送确认操作关闭服务器冲突，不要求再点击一次覆盖。冲突仍然存在时保存当前双方副本，避免同步中心展示过时的本机值。')
p('唯一值构成的字符串、数值或布尔列表，在两侧都保留原列表并只追加时，可自动合并新增值，并使用确定顺序使结果一致。存在重复值、删除、重排或其他无法证明安全的变化时不使用这条规则。用户撤销编辑也会进入待上传队列，避免被之前排队的旧内容覆盖。')

page('压缩传输与格式兼容')
h('压缩发生的位置')
p('大于 250000 字节的请求原先把原始 JSON 编码成 Base64 再分片，Base64 还会放大数据。本次在服务器声明 Gzip 能力后，对整个增量请求先压缩，再切分和编码。服务器重组分片后识别 Gzip 标记并在限定输出大小下解压，再进入既有身份校验和原子版本校验。')
p('下载时，新浏览器在分片请求中声明 encoding gzip。服务器单独压缩每个原有数据库分片，浏览器逐片解压，拼接原始 UTF8 字节，再解析 JSON 和核对最终内容哈希。每片压缩不会改变数据库分片编号；中文和 Emoji 可以跨分片字节边界，恢复后仍与原文一致。下载请求次数仍取决于原分片数，本次主要减少下载字节。')
h('实测压缩结果')
p('使用中文和 Emoji 的重复内容作为可复现测试夹具：原始 JSON 为 1000011 字节，原 Base64 为 1333348 字节，需要 17 个上传分片。Gzip 后为 2005 字节，Base64 为 2676 字节，只需要 1 个上传分片。测试中的传输体积下降约 99.8%。这是高重复内容样本，不是对真实账号所有数据的平均压缩率承诺。')
p('压缩后若体积更大，自动发送原始内容。已有压缩图片、随机内容及高熵数据通常不会获得相同收益。修改一条记录仍优先传增量，压缩只是进一步缩小需要传输的部分，不会每次强制上传全数据库。')
h('兼容与资源上限')
p('新版后台继续接受旧的原始 JSON 分片。新版客户端只有读取到 transferEncoding gzip 才压缩上传；旧后台仍走原格式。浏览器没有原生压缩或解压 API 时使用原格式。下载只在客户端明确请求时压缩，旧客户端收到的字段和分片数据保持兼容。')
p('上传原始请求及解压后请求上限均为 24000000 字节，最多 200 片；单片解压上限为 120000 字节。非法压缩内容、未知格式、校验失败、解压超限和缺失分片都不能进入数据替换流程。格式处理不使用客户端提供的用户 ID 作为授权，既有账户身份、撤销检查和服务器权限保持生效。')
h('数据库保存形式')
p('业务数据继续保存为可查询 JSON，保留版本、哈希、删除标记和历史，支持增量合并与恢复。本次没有 SQL 迁移，也不改写或清理既有用户记录。')

page('验证证据与发布顺序')
h('代码与测试范围')
p('浏览器主要改动位于 engine.ts、outbox.ts、merge.ts 与 transfer.ts；服务端新增 compression.ts，Azure sync-api.mjs 和网站的大值接口均使用同一有界解压及分片压缩规则。Functions 打包白名单和容器白名单同步加入共享模块，避免源码可以运行而生产包缺依赖。应用包和锁文件更新为 3.18.0，同步后台及 Functions 包更新为 3.4.0。')
p('新增 sync-smart.test.mts 包含 13 项回归：A 离线而 B 先同步；两侧独立编辑；撤销排队修改；持久冲突自动恢复并确认；版本回退保护；追加列表合并；中文压缩上传；跨边界下载；恶意或损坏压缩内容；真实引擎能力协商与旧后台回退；真实引擎下载哈希校验；Azure 后台压缩接口；IndexedDB 跨标签页递增编号。')
p('完整同步测试包含新增测试共 125 项；Azure 后台既有 16 项测试验证签名、账号门禁、接口可达性和错误处理。TypeScript 编译检查、修改文件 ESLint 检查及 Next.js 生产构建通过。测试使用隔离夹具，没有声称已经在真实账号的两台设备上完成生产验收。')
h('上线执行顺序')
p('本机与远程 Main 在修改前经 git fetch 核对处于同一提交。工作区中原有的界面修改被保留，本次仅提交同步代码、相关发布文件、测试和本说明。提交后再次核对远程 Main 与本机提交一致。')
p('上线时先发布 3.4.0 的 Azure Functions 后台，再发布 3.18.0 的网站前端。后台与旧页面保持格式兼容；前端在旧后台上不会提前启用压缩上传。Functions 包由既有 build azure functions 脚本产生，可按现有部署方式上传；本次没有新增云资源或修改数据库结构。')
p('提交 Main 与实际发布 Azure 后台是不同动作。本说明记录的是仓库修复和本机验证；本次没有取得 Azure 生产发布回执，因此不把压缩能力表述为已经在线启用。上线后仍应使用两个受控设备验证断网修改、恢复联网、连续同步和页面重启，并观察后台错误及实际传输体积。')
page('回退与设计参考')
h('回退与尚存边界')
p('如需回退前端，可保留新版后台，因为它仍接受原格式请求。如果回退后台，应先让客户端重新读取能力目录；未完成上传的操作始终保留在本机，重新同步后可走原格式。不应通过回滚业务数据库版本解决程序部署问题，客户端会拒绝把较旧云端版本当作最新基准。')
p('服务器无法主动识别仍处于离线设备中的未上传内容，任何系统都需要该设备重新联网后才能合并。冲突完全自动消失需要有可判断的共同基准或双方已经收敛；同字段真实分歧仍需保留选择。此次没有为文字编辑引入 CRDT，也没有变更已有备份保留策略。')
h('设计参考')
p('Microsoft Azure 文档关于乐观并发控制和 ETag 的原则：learn.microsoft.com/en-us/azure/cosmos-db/database-transactions-optimistic-concurrency。这里只参考版本校验原则，项目仍使用现有 PostgreSQL，同步实现未迁移到 Cosmos DB。')
p('MDN 原生 Compression Streams API 与 DecompressionStream：developer.mozilla.org/en-US/docs/Web/API/DecompressionStream。项目使用 Gzip 与兼容回退，不增加第三方浏览器压缩包。')

doc.core_properties.title='JackYun 同步系统升级说明'
doc.core_properties.subject='3.18.0 多设备同步与压缩传输'
doc.core_properties.author='Codex'
for element in [doc.element,*[style.element for style in doc.styles]]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
target=output/'JackYun-3.18.0-Smart-Sync-20261014.docx'
doc.save(target)
print(target)
