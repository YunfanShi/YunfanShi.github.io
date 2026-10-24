from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,1)
def page(): doc.add_page_break()

doc.add_paragraph('正式版同步故障排查与诊断补充',style='Title')
p('JACKYUN Portal 3.40.1  2026年10月24日')
p('本次工作针对正式版同步不可用及客户端报告的 quota exceeded、SYNC_UNAVAILABLE 和请求中止问题。已检查客户端日志、Vercel 运行错误、Supabase 日志与数据库元数据，以及 Azure 函数和 PostgreSQL 监控。3.40.1 补充正式环境的故障证据采集；当前证据尚不能宣告同步故障已修复。')
h('已确认的客户端与生产证据')
p('两份客户端导出日志均来自 feat/seiue-schedule-sync 分支的预览部署。新地址 aln55k05g 对应提交 64acc7e，13时43分 UTC 的两次 v3 请求返回 503 与 SYNC_CONFIGURATION_ERROR，正文明确写明 Versioned sync backend is not configured。随后 v2 上传返回 400 INVALID_OPERATIONS。Vercel 配置确认该分支没有 AZURE_SYNC_API_URL；这项事实不能解释用户另外报告的正式版故障。')
p('Vercel 生产错误记录包含 03时47分 UTC 的自动备份回执未确认，以及 05时29分 UTC 的复习队列读取失败。对应复习表的 Supabase 请求返回 200，耗时约7.7至7.9秒。Supabase 03至12时 UTC 约有308条 PostgREST timeout manager 记录，主要集中在03至05时。现有日志无法将这些超时逐一关联到同步请求。')
h('Azure 与 Supabase 资源检查')
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for c,t in zip(table.rows[0].cells,['检查项','实测结果','解释边界']): c.text=t
for row in [
 ('Azure 数据库磁盘','约13.6%','未显示磁盘耗尽'),
 ('Azure 数据库 CPU','峰值约16.9%','小时聚合不足以定位单次失败'),
 ('Azure 连接失败','当天指标为0','不能替代真实账户请求验证'),
 ('Supabase 数据库大小','约118 MB','数据库大小不是所有配额的总称'),
 ('Azure 健康接口','HTTP 200','仅证明进程存活'),
]:
    for c,t in zip(table.add_row().cells,row): c.text=t

h('已纠正的推断与日志缺口')
p('直接访问 jackyun.top 的 v3 接口返回应用 JSON 401 UNAUTHORIZED，表示请求未登录。此前 Vercel 连接器将其识别为部署保护，不能据此断言主站被 Vercel 登录墙阻断。未登录响应也不能证明真实账户同步可用。')
p('自动备份代码是在执行 create_automatic_backup_with_sync 调用后，发现错误或回执格式无效时，才抛出“自动备份尚未确认保存”。因此，不能用这条错误推断 Azure 导出读取失败，也不能把未检索到某条 HTTP 日志当作请求必未执行。')
p('Azure 应用日志、HTTP 日志和失败请求追踪均处于关闭状态；订阅的 Microsoft.Insights 为 NotRegistered，未找到应用监测配置。Vercel Hobby 原始运行日志保留范围不足以回查全天。Azure 数据库直接只读诊断连接未建立，未读取其实际操作记录。本次未更改网络权限、数据库数据或日志服务订阅。')
page()
h('诊断版本改动')
p('同步代理记录 configuration_missing、configuration_invalid、session、session_fetch、upstream_response 和 upstream_fetch 阶段。记录包含请求编号、固定接口路径、HTTP 方法、状态码、耗时、经过允许列表筛选的异常类型与系统错误码；将同一请求编号传给 Azure，以便后续跨服务对照。')
p('客户端在账户绑定初始化失败或同步循环失败时，上报异常类型、应用版本、阶段、HTTP 状态和在线状态。QuotaExceededError 可附带浏览器 origin 的 usage 与 quota 估计值；这些总量不等于 localStorage 剩余空间，仍需根据异常阶段判断。诊断发送独立于 IndexedDB 和 localStorage，每个页面最多每分钟一次。')
p('新增服务端接口要求已验证的登录用户，拒绝跨站 origin，限制输入为2048字节，并按用户在进程内限制每分钟一次。服务端重新筛选字段；不接受任意错误正文、堆栈、学习键值、Cookie、令牌或 API 密钥。该进程内限制是减噪措施，不宣称跨实例全局限流。')
p('诊断传输失败被隔离，不改动同步结果、待上传队列、回执、冲突内容或恢复基线。本次不因尚未确定的 quota 原因自动删除任何学习数据。')

h('验证与发布范围')
p('验证包含诊断输入中敏感字段的排除、非法错误类型与状态的拒绝、DOM QuotaExceededError 的保留、认证与跨站检查、无 Content Length 情况下的流式大小限制，以及代理网络异常日志不包含私密错误正文。同步回归覆盖原有重试、冲突、不可变发送、回执恢复和分片下载。')
p('本地 Next.js 生产构建及 TypeScript 检查通过，完整同步回归168项通过，涉及文件的 ESLint 与私钥标记扫描通过。测试和构建通过仅说明代码层验证结果，不代表已完成真实账户的跨设备同步验收。实际发布状态以提交后的 Vercel 生产部署与主站响应检查为准。')
h('下一次失败的判断方式')
p('出现 sync-diagnostic 且 errorName 为 QuotaExceededError，可确认客户端存储写入异常；initialization 表示账户绑定或设备设置阶段，cycle 表示同步循环阶段。此时继续保留队列，排查该 origin 的学习数据与可再生缓存占用，不以清站点数据作为恢复步骤。')
p('出现 sync-proxy 的 session 或 session_fetch，应追查会话刷新与认证服务；upstream_fetch 应结合 TimeoutError、AbortError 或系统连接码判断网络失败；upstream_response 表示 Azure 确实返回了错误状态，可凭 upstreamRequestId 进一步定位服务端请求。')
p('若新诊断版本没有收到客户端报告，不能推断故障消失；可能是客户端尚未加载3.40.1、未在线、诊断请求被认证拒绝，或失败发生在未覆盖的其他页面。现有页面需正常加载新版本后，客户端诊断才开始生效；新的主站代理日志随服务器部署生效。')
h('后续验收与恢复')
p('后续验收需使用正式账户，在两台客户端完成一次编辑、上传确认、另一端拉取和刷新恢复，并确认待上传数量归零与内容一致。本次不把健康接口200、未登录401或部署READY视为该验收已通过。若诊断改动引入问题，可回退对应提交；应用数据与数据库结构未因诊断改动变化。')
p('文档视觉校验如使用 docx-preview 浏览器回退，将明确标注该方式，不将其表述为 LibreOffice 分页验证。原始客户端日志和受保护数据库连接日志仅保留在本机，不上传仓库。')

footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('JACKYUN 3.40.1  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for paragraph in list(doc.paragraphs)+[p for row in table.rows for cell in row.cells for p in cell.paragraphs]:
    for run in paragraph.runs:
        run.font.name='Microsoft YaHei'
        run.font.color.rgb=RGBColor(0,0,0)
        run.font.underline=False
        fonts=run._element.get_or_add_rPr().rFonts
        fonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+attr),None)
        color=run._element.get_or_add_rPr().find(qn('w:color'))
        for attr in ['themeColor','themeTint','themeShade']:
            color.attrib.pop(qn('w:'+attr),None)
for cell in table.rows[0].cells:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
doc.save(out/'正式版同步故障排查与诊断补充3.40.1.docx')
print('DOCX created')
