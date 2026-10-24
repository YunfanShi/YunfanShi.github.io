from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

out = Path(__file__).parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(15)
def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text, 1)

doc.add_paragraph('正式版同步账号核验恢复修复', style='Title')
p('JACKYUN Portal 3.40.2 / Azure Sync 3.6.1   2026年10月24日')
h('一、用户反馈与本次范围')
p('用户确认3.40.1正式版刷新、重新同步后仍然失败。此前3.40.1仅补充诊断，不是恢复修复。本次根据新增正式版客户端日志及生产代理日志，修复账号权威查询短暂不可用造成的同步拒绝，并改善网站服务端只读请求的恢复能力。保留全部本地内容、待上传操作、冲突和回执。')
h('二、已经关联的正式请求')
p('北京时间22:16:57（UTC 14:16:57），jackyun.top 的 GET /api/sync/v3 返回503，错误为 ACCOUNT_STATE_UNAVAILABLE / JY-5002。Azure请求编号为1f130bf3-3090-49ff-8777-03569f383405，网站代理编号为4559da4f-ca8a-46d5-833f-8f089fae0dda。代理明确记录 upstream_response，耗时8351毫秒，客户端记录8718毫秒。说明请求已到Azure，由后端账号状态门禁拒绝。')
p('同一时段，网站授权查询报 Authorization lookup failed，离线工作区读取复习队列失败；通知快照耗时约10秒，自动备份未取得确认。这些错误说明账号及业务读取链路同时受影响；不能把所有失败都归因于浏览器配额。新日志未提供可以确认 QuotaExceededError 的本机写入证据。')
h('三、数据库和链路证据')
p('严格限定UTC 14:16:40至14:17:40后，Supabase profiles 接口多个HTTP 200请求的 origin_time 达6716、6887、7482、7776、7823、7861毫秒。原Azure账号核验每次只等5秒，首次核验没有恢复重试。链路延迟足以触发该门禁不可用。日志没有跨云统一追踪，因此不声称某一条Supabase请求已和该Azure编号逐一匹配。')
p('pg_stat_statements累计统计：同步用profiles三列查询27456次，平均执行0.53毫秒，最大240.76毫秒；网站用role等三列查询34305次，平均0.84毫秒。累计SQL统计不能代表单次完整网络耗时，但没有支持该简单查询持续执行7至8秒的证据。本次未盲目新增索引、修改RLS、升级套餐或迁移账号权威源。')

doc.add_page_break()
h('四、修复内容与安全边界')
p('Azure账号权威读取的每次截止时间调整为8秒；临时不可用最多再读一次，最多两次。对于相同已验证用户、session_id和签名有效期的并发核验，合并在途请求，并将同时在途会话数量限制在500以内。冷实例和首次会话仍必须从权威源读到active状态。')
p('HTTP 401或403、账号缺失、停用、删除、身份不匹配直接拒绝。签名和声明核验仍先执行。原有成功会话缓存及其有效期边界保留，不允许使用其他账号、可编辑metadata或未经验证的Cookie恢复授权。缓存不能替代即时账号生命周期迁移。')
p('网站服务端和中间件的PostgREST GET/HEAD请求，在网络失败、超时或指定临时HTTP状态时最多再试一次。调用方已经取消则不重试；429带Retry-After时直接交回调用方。Auth保留SDK原恢复策略，POST/PATCH/DELETE等写操作不自动重放，避免重复修改和重复计数。读取重试最多增加一次原请求预算，不保证上游持续故障时一定恢复。')
h('五、版本、构建与部署记录')
p('任务开始时，隔离工作区和远程Main同为790523149733ec0f7a3395c4e37e578861e344e4，网站版本3.40.1。修复网站递增为3.40.2，Azure同步包由3.6.0递增为3.6.1，并更新运行健康版本。原始工作目录存在大量未提交改动，继续保留；没有通过清理、覆盖或重置方式强制同步。')
p('Next.js生产构建及TypeScript检查通过，涉及的Supabase客户端与恢复模块ESLint通过。Azure Functions依照锁文件构建355个文件，压缩包562985字节，SHA-256：3d1ecbe9f4e2ca61c3bb596edad1d06484eeacf455b840c4b5300e69f5770981。本轮未新增或运行自动化测试；构建检查不能替代真实账号多设备验收。')
p('部署及主站核对结果以同目录release-receipt.json为准。后端需要实际运行3.6.1，网站需要Main对应Vercel生产部署READY且jackyun.top加载3.40.2。健康接口只确认后端版本和进程，不能据此宣告账号同步成功。')
h('六、验收、限制和回退')
p('已定位并修复本次正式请求中的账号查询恢复缺口。Supabase为何在该时段出现数秒链路延迟仍未由现有日志证明，本次没有将其描述为已消除的云服务根因。真实账户的编辑、上传确认、另一客户端拉取和刷新恢复仍需通过实际会话观察；未取得这些证据前不宣称端到端恢复已验收。')
p('若修复引发问题，可回退本次提交并重新部署对应后端包，无数据库结构迁移。发布后继续检查生产upstream_response及账号状态不可用是否复发。不要清除站点数据；这会损失诊断现场并可能影响尚未确认保存的内容。文档采用docx-preview浏览器回退渲染核对，未冒充LibreOffice分页验证。')
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Microsoft YaHei'
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.underline = False
        props = run._element.get_or_add_rPr()
        fonts = props.rFonts
        fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + attr), None)
        color = props.find(qn('w:color'))
        for attr in ['themeColor', 'themeTint', 'themeShade']:
            color.attrib.pop(qn('w:' + attr), None)
doc.save(out / '正式版同步账号核验恢复修复3.40.2.docx')
print('DOCX created')
