from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out=Path(__file__).parent
doc=Document()
section=doc.sections[0]
section.page_width,section.page_height=Inches(8.5),Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1']:
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(15)
def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,1)
def page():doc.add_page_break()

doc.add_paragraph('同步系统可靠性改造实施说明',style='Title')
p('JACKYUN 网站 3.44.0  Azure 同步后端 3.7.0  2026年10月25日')
p('本报告说明本地保存、离线业务幂等、统一恢复调度和v4增量同步的实现及验证。代码保留Azure与Supabase架构、已有回执、三方合并和分片校验。发布采用能力开关；Azure迁移未完成时继续使用v3，不把已编写的v4代码当作生产增量同步已启用。')
h('版本基线与工作区')
p('任务开始时本地实施工作区为4c55462，远程Main已前进至fd0a8d2，网站版本3.43.0，包含SEIUE只读同步功能。已安全快进干净工作区并保留该功能；实施期间Main另有3.43.1和3.43.2修复，将一并合入发布。根据新模块递增规则，本次网站版本使用3.44.0，避免按旧计划发布3.41.0造成倒退。Watch源码不在本次修改范围。')
h('本次实际变化')
p('同步白名单数据写入统一Repository，IndexedDB文档与上传意图在同一事务中提交；localStorage作为兼容投影。页面可以立即显示草稿，但“已保存到本机”由事务完成事件决定。保存失败会保留内存草稿并提供恢复文件；恢复文件同时包含草稿与兼容投影，不宣称浏览器强制终止前尚未完成的写入已经持久化。')
p('工作区、任务复习业务与自动备份共用恢复策略，分别运行在独立队列。新增有抖动的退避、依赖熔断、跨标签页锁与租约、稳定操作编号和业务回执。备份失败不等待主同步完成，也不会清除备份改动标记。')
h('发布边界')
p('Supabase业务幂等迁移及Azure变化日志迁移均已应用。Azure迁移先因连接默认只读而安全回滚，随后仅在迁移会话显式开启写事务并完成。临时单IP网络规则经用户明确授权，执行后删除并核对。运行角色、旧v3写入日志、重复抑制、CAS、删除标记、账号序号和隔离均通过回滚式验证。v4仍关闭，等待真实账号灰度验收。')
p('网站3.44.0已通过Vercel生产部署并绑定jackyun.top，源码提交199fde9；Azure健康接口确认3.7.0。测试、构建和部署READY不等于真实账号两客户端验收完成，尚待登录会话完成该项验证；详细证据见verification.json。')

page()
h('本地持久化与账号隔离')
p('新仓库按账号划分命名空间。旧仓库保留原状，迁移复制队列、元数据、冲突和设置，并对复制后的记录计算哈希、核对数量；只有读回校验通过才写入迁移完成标记。中断后可重新执行，不使用清站点数据作为恢复方法。')
p('文档写入同时分配递增操作序号，保留最早云端基线并合并尚未发送的后续编辑。已经发送的请求正文单独冻结；新的编辑不会改变其操作编号或正文。云端应用使用本机内容比较检查，发现更新的本机编辑即推迟覆盖。')
p('跨标签页编辑按原投影、本机新值和当前已保存文档做三方合并。独立字段可以合并，同字段矛盾拒绝写入，原已保存内容、队列和未保存草稿仍保留。设备时间不会决定学习内容胜负。')
h('旧页面和账号切换恢复')
p('Repository记录最后成功投影的内容。旧iframe或原生localStorage写入如果发生在浏览器退出前，重新打开时会与该投影基线核对，恢复尚未进入新仓库的编辑，而不是直接用旧投影覆盖新数据库。')
p('账号切换先把旧账号最后的原生编辑归还其仓库，再恢复目标账号投影。切换本身清空的投影不被解释为学习数据删除。异步同步Worker固定绑定账号数据库，避免请求结束时误写入新账号。跨标签页账号切换还会停止旧上下文保存；上传携带预期账号，服务端发现签名身份不同会拒绝。')
h('空间不足与本机恢复副本')
p('遇到QuotaExceededError只清理可重新下载的分片缓存，并且最多重新尝试一次；不会删除文档、待上传操作、冲突、冻结请求或恢复基线。兼容投影写入失败与权威数据库保存失败分开提示。数据库保存失败时不宣布保存成功，恢复文件可下载当前内存草稿和已有投影。')
p('网页暂停、切后台和退出不能可靠等待任意网络请求，因此保存以平时的事务提交为准；休眠和主动退出继续等待已开始的本地保存及备份。浏览器撤销整个站点存储、设备损坏和用户删除恢复文件不属于可由网页同步代码消除的风险。')

page()
h('幂等回执与数据库事务')
p('离线业务操作先进入持久化队列，再尝试网络发送。请求包含稳定operationId及根据kind和payload计算的哈希。服务器重新计算哈希；数据库比较操作类型、完整payload和请求哈希，相同操作返回原回执，同编号内容变化拒绝。回执缺失或不匹配时不会清除队列。')
p('复习操作的记录新增、间隔与熟练度计算、复习项目更新以及回执新增在一个数据库事务中完成。任务和计划操作也通过同一受限RPC处理。本地任务与复习视图、上传意图在同一IndexedDB事务保存；新建任务的本机编号直接使用操作编号，避免离线后继续勾选时引用临时占位编号。客户端不能直接插入回执，匿名角色不能读取。公开RPC是SECURITY INVOKER包装器；真正需要写入受保护回执的函数放在未公开schema内，验证当前身份及active账号，并对所有业务读写使用当前用户条件。')
h('统一调度与鉴权边界')
p('失败退避起始约2秒，上限2分钟，并加入随机抖动；服务端Retry-After沿代理传递并被遵守。连续5轮失败后熔断30秒，只允许一个恢复探测。恢复状态写入账号仓库，手动重试仍受到跨标签页互斥约束。')
p('Web Locks可用时优先使用；不可用时使用IndexedDB原子租约、递增令牌与定时续租。过期Worker停止本机云端应用。SDK会话恢复只针对明确401进行一次受控重试；失败后暂停上传，登录恢复事件唤醒队列。403权限拒绝不通过扩大缓存有效期放行。Supabase仍是账号状态权威源。')
h('v4增量协议与兼容')
p('新增/api/sync/v4能力协商、上传、分页变化读取及回执查询。变化日志由legacy_sync_data触发器记录，v3和v4写入都会进入同一日志；账号序号在同一事务内递增。活跃账号写入时每日清理超过30天的变化日志，删除标记和幂等回执首版不自动清理。')
p('首次建立基线时先读取账号序号，再扫描基线并追平变化。每页先持久暂存，全部应用后才原子推进检查点并清除暂存。单键失败会保留第一个未完成页和检查点；之后页面仍可读取和应用，下一轮借助元数据跳过已经成功的记录。过期游标触发一次基线重建，不清空本机待上传数据。')
p('v4只在服务端明确声明可用时启用。未启用或数据库依赖未准备好时继续v3；已经发送且结果未知的请求不切换协议重发。大值继续使用现有分片通道及版本哈希绑定，避免在本轮同时替换全部传输机制。')

page()
h('自动化与浏览器验证')
table=doc.add_table(rows=1,cols=3)
table.style='Table Grid'
for cell,text in zip(table.rows[0].cells,['检查项目','结果','证据边界']):cell.text=text
for row in [
 ('同步回归','182项通过','含新持久化与v4恢复场景'),
 ('鉴权恢复回归','52项通过','包含会话失败和退出保留数据'),
 ('Azure后端回归','25项通过','含v4能力门禁及回执拒绝'),
 ('Edge浏览器验证','7个场景通过','真实IndexedDB及标签页 合成身份'),
 ('私钥标记扫描','3项通过','不代替所有安全审计'),
 ('Supabase契约','通过并回滚','任务和复习重复请求不重复生效'),
 ('生产构建与类型','通过','最终发布仍需核对线上版本'),
]:
 for cell,text in zip(table.add_row().cells,row):cell.text=text
p('浏览器验证覆盖两标签页独立字段合并、刷新后文档及待上传意图恢复、原生旧页面写入后的重启恢复、账号投影隔离和原账号恢复，以及业务快照与意图原子提交和配额失败回滚。它使用合成身份及本机服务，不是正式账号跨设备验收。')
p('合并后扩展学习回归185项通过、3项失败；失败涉及既有CUE偏好字段预期、游戏面板Hook夹具和移动工具栏aria-controls断言。核对Main发现相同实现或缺失已存在，本次未把这些失败描述为同步回归通过，也未扩大修改范围。')
h('数据库迁移和上线步骤')
p('Supabase迁移记录为20261025023755 reliable_offline_mutations。数据库已核对客户端无回执插入权限、匿名无读取及RPC执行权限。任务幂等与哈希冲突检查、复习幂等和调度更新检查均使用回滚事务；复习表没有现成测试项目，因此创建了事务内临时项目并全部回滚。')
p('Azure SQL位于azure/sync-v4.sql。管理员执行前需验证目标库、备份迁移前结构，并使用ON_ERROR_STOP及事务执行。迁移完成后，以运行身份检查表权限、触发器、序号和回执访问；确认旧v3写入会记入日志，再按账号启用SYNC_V4_USERS，验收后才启用SYNC_V4_ENABLED。不得在迁移失败时打开能力开关。')
h('回退与待完成验收')
p('应用回退保留新增数据表和队列，不执行破坏性降级或清库。v4初次启用前可以关闭能力开关继续v3；如果已有结果未知的v4发送记录，必须保留并恢复该请求，不可直接改成另一协议。自动备份使用独立恢复状态，其503可通过经过筛选的错误分类和统一请求编号追查，不输出学习正文、数据库消息或登录凭据。')
p('Azure生产变化日志迁移与运行身份验证已完成，验证SQL保存在azure/tests/sync-v4.sql，所有夹具写入均回滚。待完成事项为真实账号v4灰度启用，以及两客户端编辑 上传确认 拉取 刷新和离线恢复。临时规则只允许当前单个出口IP，结束后删除并核对；未配置全量v4开关。')
p('设计参考 Replicache How Replicache Works与Push Endpoint Reference，以及Microsoft Azure Circuit Breaker模式。详细来源链接保存在同目录README。文档使用DOCX渲染器检查；若本机无LibreOffice则使用明确标注的docx-preview回退。')

for style in doc.styles:
 for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
props=table._tbl.tblPr
borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
props.append(borders)
for cell in table.rows[0].cells:
 shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
paragraphs=list(doc.paragraphs)+[p for row in table.rows for cell in row.cells for p in cell.paragraphs]
for para in paragraphs:
 for run in para.runs:
  run.font.name='Microsoft YaHei';run.font.color.rgb=RGBColor(0,0,0);run.font.underline=False
  props=run._element.get_or_add_rPr();props.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
  for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:props.rFonts.attrib.pop(qn('w:'+attr),None)
  color=props.find(qn('w:color'))
  for attr in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+attr),None)
doc.save(out/'JACKYUN同步可靠性改造实施说明3.44.0.docx')
print('DOCX created')
