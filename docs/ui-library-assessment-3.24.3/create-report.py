from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, subprocess
root=Path(__file__).resolve().parents[2]
out=Path(__file__).parent
D=Document(); sec=D.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.9); sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=D.styles[name]; s.font.name='Microsoft YaHei'; s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 s.font.size=Pt(10.5 if name=='Normal' else 23 if name=='Title' else 15 if name=='Heading 1' else 12)
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.16
D.styles['Normal'].paragraph_format.widow_control=True
pages=[]
def p(t): D.add_paragraph(t)
def h(t): D.add_heading(t,level=1)
def page(): D.add_page_break()
D.add_heading('JackYun 点击响应与界面库更换评估',0)
p('版本 3.24.3  |  2026 年 10 月 15 日')
p('结论：可以更换或引入组件库，但不建议先重写前端框架。优先保留 Next.js 和 React，统一交互组件，再改善页面预取、数据等待和重组件渲染。这样更直接对应“别人点一下就出来，而这里不好用”的问题。')
h('用户关注的问题')
p('本次关注点击后内容能否迅速出现，以及操作是否连贯。外观平面与响应迟缓可能同时存在，但不能只靠增加阴影或动画解决。评估需要分别衡量按下后的即时反馈、弹窗或标签页的本地切换，以及需要网络数据的页面打开。')
h('实际使用的技术')
p('package.json 声明 Next.js 16.3.8、React 和 React DOM 19.2.8、Tailwind CSS 3.4.17，Node.js 要求至少 22。Node.js 是运行环境；React 负责组件，Next.js 负责路由和服务端渲染，Tailwind 负责样式。项目没有声明一套覆盖按钮、弹窗和抽屉的通用组件库。')
p('现有界面包含手写 React 组件和通过 iframe 嵌入的旧 HTML 页面。因此，只换外层组件库不会自动改善 iframe 内部的按钮与脚本。已有的拖拽工具和 Markdown 渲染库也不是完整的交互组件系统。')
h('为什么同样做 Web 别人的更顺')
p('常见实现会立即显示本地面板、保留上一份数据，在后台读取新数据；会提前准备常用页面，并保持选中、忙碌、成功和失败状态一致。这些体验主要取决于组件实现、状态组织和数据链路，而不是 Node.js 的名字。')
p('本次完成源码审阅与官方资料核对，尚未取得用户实际设备的 Performance trace、线上 INP 或对比网站的测试数据。以下结论分为已确认实现与性能假设，不声明已测得任何加速倍数。')
page(); h('源码证据和影响')
for a,b in [
('导航预取覆盖不足','src/components/layout/sidebar.tsx 第 29 至 39 行和第 166 行：Link 使用 prefetch={false}，手动 prepare 只绑定鼠标进入和焦点。触屏点击前没有保证有足够预取时间。首页 product-card.tsx 第 25 至 28 行也关闭预取，未添加同类手动预取。首次打开更容易等待路由请求；这是已确认的策略，实际等待时长待测。'),
('反馈以悬停为主','product-card.tsx 以 hover 位移、阴影和图标缩放为主，没有组件自身的 active 按压或 pending 状态。sidebar.tsx 导航有加载提示，并非完全没有反馈，但高亮依赖当前 pathname。触屏无法完整获得桌面 hover 手感。'),
('数据在主体返回前等待','src/app/(portal)/schedule/page.tsx 先 await 身份读取，再按登录状态读取角色，之后才返回包裹 ScheduleControl 的 Suspense；该内层边界不能提前展示它之前的 await。父级 loading.tsx 已存在。dashboard/page.tsx 同样先读取用户，再等待八个并行统计请求，页面内容受最慢请求影响。'),
('复杂页面需要单独定位','schedule-control.tsx 同时承担多类学习日程与 Cue 交互，有多处本地状态与派生数据。代码体积和状态数量只说明值得做 Profiler，不等于每次点击都整页重绘。主控制台定时器是 60 秒，不能据此指认每秒刷新导致卡顿。'),
('旧页面加载有独立成本','legacy-frame.tsx 缓存 HTML fetch，但挂载时仍需替换和注入 HTML，再通过 srcDoc 建立 iframe 文档。缓存 HTML 不等于保留 iframe 运行状态；路由重挂载、依赖变化与脚本初始化的实际成本需要录制确认。'),
('动画可能被主动关闭','globals.css 支持系统 reduced-motion、html[data-reduced-motion] 和页面 motion=false。当前样式也移除了原生触摸点击高亮。需要检查真实设备设置，再判断是否应补充自定义反馈，不能简单强制开启所有动画。')]:
 D.add_heading(a,2); p(b)
page(); h('更换方案比较')
t=D.add_table(rows=1, cols=3); t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['方案','能改善什么','判断']): c.text=txt
for row in [
('保留现有框架并统一基础组件','按钮按压、选中、局部加载；同时调整导航和数据读取','首选，现有业务复用最多'),
('渐进引入 Radix Primitives','弹窗、菜单、标签页的焦点与键盘行为；样式自行统一','适合自定义设计，先做小范围样板'),
('引入 Motion for React','弹窗、列表和布局变化的连续过渡','按需补充，不能消除网络等待'),
('整体替换为成熟成套 UI','快速统一复杂表单和控件的外观及行为','可行，但要适配主题、触屏与现有页面'),
('迁移到 Vue 或 Svelte 等框架','改变组件与渲染模型','成本高，尚无证据支持现在重写')]:
 for c,txt in zip(t.add_row().cells,row): c.text=txt
p('以上是工程判断，不是库性能排名。Radix 提供无样式交互基础；Motion 提供动画能力，二者职责不同。若采用成套 UI，仍需在真实业务页验证依赖体积、样式覆盖和触屏操作，不能凭演示站点判断收益。')
h('建议的落地顺序')
p('第一步：选侧栏入口、一个弹窗、一个标签页和日程中的一个常用操作作为样板。按钮按下立即反馈；弹窗先显示框架再取数据；切换标签时只更新必要区域；保存时允许本地先反馈，但失败可恢复并明确提示。')
p('第二步：按使用频率和网络条件有限预取常用入口，避免一次预取全部模块。为需要等待的数据设置更细的加载边界，优先呈现主体，并行读取互不依赖的数据。权限判断仍由服务端执行，不为了速度绕过授权。')
p('第三步：依据 Performance 和 React Profiler 结果拆分重组件、缩小状态更新范围，必要时虚拟化长列表。检查同步存储、JSON 序列化、旧页面脚本初始化及动画布局成本，确认后再优化。')
p('第四步：样板若证明交互原语缺失是主要维护成本，再渐进引入 Radix 等基础组件；确实需要弹性或共享布局过渡时再加入 Motion。无需先更换 React 或 Node.js。')
page(); h('验证方式与验收目标')
p('必须使用生产构建或生产站点，避免把 next dev 的首次编译误认为线上点击性能。固定同一设备、同一账号和数据量，分别测试冷首次打开、二次打开、触屏与鼠标操作、正常网络与慢网络。匿名测试不能代表登录后的统计与权限链路。')
p('每类常用操作至少重复 20 次，记录点击到首次反馈、点击到可用内容、路由请求等待和主线程长任务，比较中位数与 P95。记录 60 Hz 设备上动画是否持续超出约 16.7 ms 的帧预算；120 Hz 设备的预算更短。')
p('建议目标：本地按钮按下反馈尽量在 100 ms 内；无需网络的小弹窗或标签切换尽量在 200 ms 内出现可用内容。网络页面应立即显示选中或忙碌状态，完整内容的目标应根据基线另定。以上属于建议目标，不是本次实测结果。')
p('Google 对良好 INP 的建议阈值为 200 ms 或以内，正式评估看真实访问分布的第 75 百分位。INP 只衡量交互到下一次绘制，不能替代“完整页面数据已经可用”的耗时，因此两个指标需分别观察。')
h('本次交付和仓库范围')
p('开始时成功读取远程 main，本机 HEAD 与 origin/main 均为 1dfe2588bad9cd5e9956c8f246dc77c7d14e1e0d，提交差异为 0 和 0。工作区存在原有未提交内容，因此仅对比提交一致性，不能称工作区所有文件与远程完全一致。')
p('本次只交付评估文档、可复现文档生成源和版本元数据。按用户规则将补丁版本从 3.24.2 升为 3.24.3；未实施界面库替换，也未声明修复完成。只提交本次文件及版本字段，保留已有源码和 package.json 中其他未提交改动。最终提交和推送结果以交付消息及 Git 记录为准。')
h('官方资料')
for txt,url in [
('Next.js 导航与预取','https://nextjs.org/docs/app/getting-started/linking-and-navigating'),
('Google 优化 INP','https://web.dev/articles/optimize-inp'),
('Google INP 定义与评价','https://web.dev/articles/inp'),
('Radix Primitives 简介','https://www.radix-ui.com/primitives/docs/overview/introduction'),
('Motion for React','https://motion.dev/docs/react')]:
 p(txt+'  '+url)
# Remove template theme colors and decorative borders.
for style in D.styles:
 for el in list(style.element.iter(qn('w:color'))):
  el.attrib.clear(); el.set(qn('w:val'),'000000')
 for el in list(style.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for row_index,row in enumerate(t.rows):
 for cell in row.cells:
  tcPr=cell._tc.get_or_add_tcPr()
  margins=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   el=OxmlElement('w:'+side); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
  tcPr.append(margins)
  for el in list(tcPr.iter(qn('w:shd'))): el.getparent().remove(el)
  if row_index==0:
   el=OxmlElement('w:shd'); el.set(qn('w:fill'),'EEEEEE'); tcPr.append(el)
  for para in cell.paragraphs:
   para.paragraph_format.space_after=Pt(4); para.paragraph_format.space_before=Pt(4)
   para.paragraph_format.line_spacing=1.15
   for run in para.runs:
    run.font.color.rgb=RGBColor(0,0,0); run.font.size=Pt(10); run.bold=row_index==0
D.save(out/'点击响应与界面库更换评估3.24.3.docx')
print('saved',out/'点击响应与界面库更换评估3.24.3.docx')
