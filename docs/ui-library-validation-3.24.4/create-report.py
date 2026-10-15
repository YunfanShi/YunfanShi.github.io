from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
out=Path(__file__).parent
D=Document(); sec=D.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.75); sec.left_margin=sec.right_margin=Cm(1.9)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=D.styles[name]; s.font.name='Microsoft YaHei'; s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 s.font.size=Pt(10.5 if name=='Normal' else 22 if name=='Title' else 15 if name=='Heading 1' else 12)
 s.paragraph_format.space_after=Pt(7); s.paragraph_format.line_spacing=1.16
D.styles['Normal'].paragraph_format.widow_control=True

def p(x): D.add_paragraph(x)
def h(x): D.add_heading(x,1)
def sub(x): D.add_heading(x,2)
def page(): D.add_page_break()
def table(headers, rows):
 t=D.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row): c.text=x
 for i,row in enumerate(t.rows):
  for c in row.cells:
   pr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side); el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
   pr.append(margins)
   if i==0:
    el=OxmlElement('w:shd');el.set(qn('w:fill'),'EEEEEE');pr.append(el)
   for para in c.paragraphs:
    para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.1
    for r in para.runs: r.font.size=Pt(10);r.bold=i==0
 return t

D.add_heading('JackYun 组件库更换计划与可行性报告',0)
p('2026 年 10 月 15 日  |  文档交付版本 3.24.4')
p('建议执行渐进式组件迁移，首选 shadcn/ui 组件源码配合 Radix 交互基础，保留 Next.js 与 React。项目已有 React、Tailwind 和 CSS 主题变量，具备接入条件；主要风险集中在主题适配、旧页面和复杂业务交互，而非必须重写整站。')
h('决策结论')
p('技术接入可行性较高，迁移成本可控，但“换库后一定更快”尚不能成立。组件库可以统一按钮、弹窗、菜单和标签页的行为；页面等待网络数据、同步脚本阻塞以及大组件更新，需要在同一计划里分别处理。')
p('推荐先做常用操作的第一批迁移，再扩展到主要 React 页面。不要一开始替换所有控件，也不要只更换圆角、阴影和配色。交付目标是点击立即有反馈、面板快速打开、任务不中断和状态一致。')
h('范围与投入')
p('第一批包括基础按钮与输入、弹窗和抽屉、标签页，以及首页入口、侧栏和工具箱的高频操作。按一名熟悉项目的前端工程师估算，第一批约 7 至 12 个工作日；主要 React 界面约 15 至 25 个工作日。旧 HTML 业务模块另估 8 至 20 个工作日以上，需先盘点再承诺。上述为规划估算，不是测试结论或交付承诺。')
h('可行性证据与边界')
p('本次以源码审阅和官方文档核对支持计划判断。当前 src/components 下共有 195 个文件；用弹窗语义或全屏定位特征筛出 29 个候选文件，仅供估算，不能等同于 29 个真实弹窗或迁移组件数。')
p('本次按用户最新要求只提交计划与报告，不交付迁移实现。尚未完成安装兼容性验证、生产构建、业务回归或真实设备性能测试。精确依赖版本和性能收益须在实施第一阶段确认。')

page();h('组件库选型')
table(['候选方案','项目适配判断','代价与建议'],[
('shadcn/ui 加 Radix','与现有 Tailwind 和自定义主题接近，适合逐个迁移','首选。需要维护纳入仓库的组件源码及统一样式'),
('Ant Design','适合表格、筛选、复杂表单密集的管理页面','备选。需要统一设计语言并处理样式顺序和服务端样式接入'),
('MUI','适合希望采用一致 Material 风格的整套界面','备选。需增加 App Router 样式集成并重做现有主题映射'),
('只用 Radix Primitives','可统一弹窗与菜单等行为，样式高度自由','可行但基础视觉仍需大量自建，不如首选方案直接'),
('更换为其他前端框架','改变组件体系，业务层迁移范围大','暂不建议，缺少必须重写 React 的依据')])
sub('为什么推荐 shadcn/ui')
p('shadcn/ui 提供可放进项目自行维护的组件源码和分发方式，严格说并非传统的封闭 npm 组件库。它适合把现有散落的界面实现收敛为 src/components/ui 下的统一组件，并保留产品自己的外观。Radix 负责所选组件的交互基础。')
p('本计划统一选择 Radix 路线，不在首批混入多套交互原语。按钮、卡片等基础元素不需要全部依赖 Radix；动画先采用少量 CSS transform 与 opacity，不把 Motion 或大型动效系统作为前置条件。')
sub('选型的适用条件')
p('首选方案适用于继续保留当前学习门户的轻量设计和渐进迁移。如果目标改为整站采用标准后台风格，且表格、筛选和表单占主导，Ant Design 的优先级可以提高。MUI 则适合明确选择 Material 视觉规范的情况。')
p('官方文档支持 Next.js 接入和渐进采用，因此架构上存在明确路径；这不代表已经验证当前精确版本组合、全部旧样式或每个业务页面。')

page();h('兼容性与迁移前置条件')
table(['现状或风险','接入处理','实施阶段通过标准'],[
('Next.js 16.3.8 与 React 19.2.8','仅交互基础组件标记 use client，保留服务端页面边界','生产构建通过，无新增 hydration 或 hooks 错误'),
('Tailwind CSS 3.4.17','优先保留 v3，人工核对生成源码；不能直接套用只适用 v4 的样式','选定组件在 v3 下完整生成样式，无缺失 utility'),
('data-theme 与 darkMode media 并存','把 light gray dark 三种主题映射到同一套 token 与选择器','三主题下按钮、浮层和输入状态一致'),
('主题变量已有十六进制颜色','映射到组件 token，避免直接套 hsl(var()) 等不兼容形式','常态、悬停、选中、禁用、聚焦颜色均正常'),
('Dialog Portal 通常挂到 body','主题 token 放在根级，统一浮层层级；核对全屏容器','全屏、手机键盘和多层弹窗不丢失或遮挡'),
('旧 HTML 使用 iframe 和注入桥接','新库只覆盖 React 区域，旧模块逐个迁移','iframe 内外的存储、消息和焦点链路保持有效')])
sub('不把 Tailwind 升级绑定到换库')
p('shadcn 当前文档覆盖 Tailwind 4 与 React 19，同时说明已有 Tailwind 3 应用可保留原配置。项目实际为 Tailwind 3 配合 React 19，因此需要逐项检查生成源码和依赖，而不能把两个独立兼容声明当作精确组合已通过。若遇 v4 专属能力，先做兼容适配，必要时改用直接封装 Radix；Tailwind 大升级单独评估。')
sub('数据与权限边界')
p('首批替换展示与交互层，沿用现有服务端动作及数据契约。弹窗可先呈现外壳，权限内容仍在授权确认后展示；保存动作可提供忙碌和可恢复状态，不能用成功动画替代真实保存结果。')

page();h('分阶段实施计划')
table(['阶段与预计投入','具体工作','交付与放行条件'],[
('阶段零  1 日','锁定页面范围、依赖路线和样式规则，录制当前高频操作基线','迁移清单、版本锁定方案和基线记录'),
('阶段一  2 至 3 日','建立 Button Input Dialog Sheet Tabs Select Skeleton 等基础组件，适配三主题','独立样板与构建检查通过，基础行为统一'),
('阶段二  3 至 5 日','迁移工具箱、常见弹窗、首页入口和导航反馈','形成第一批可体验版本，用户完成真实操作对比'),
('阶段三  4 至 7 日','迁移日程、Cue、设置的高频控件，按需拆开重面板','学习、计时、保存和快捷键回归通过'),
('阶段四  2 至 4 日','清理重复样式、核对依赖体积并完成移动端和离线回归','主要 React 界面统一，遗留项有明确清单')])
p('上述合计 12 至 20 人日，加约 20% 的集成缓冲，按 15 至 25 个工作日安排主要 React 界面。第一批为阶段零至二，计 6 至 9 人日，含缓冲约 7 至 12 个工作日。估算假设单人熟悉项目、设计方向稳定、外部服务无新增改造；实施期间出现新需求应调整范围和日期。')
sub('先迁哪些真实位置')
p('tools/tools-tabs.tsx 先统一标签切换；layout/sidebar.tsx 与 product-card.tsx 统一导航反馈；support-conversation-dialog.tsx 等常用浮层统一打开、关闭与焦点行为。之后再进入 settings 和 schedule 的复杂流程。')
sub('旧 HTML 页面另开迁移批次')
p('MockPortal、TimetableHub 等旧页面涉及内嵌脚本和存储桥接，应按用户任务逐个拆出 React 界面，保留数据兼容与恢复路径。不要用全局替换按钮标签的方式迁移。所需 8 至 20 个工作日以上为粗估，不包含在主要 React 界面的 15 至 25 日内。')

page();h('把使用舒服写成具体交付要求')
table(['用户操作','目标体验','实现约束'],[
('按按钮或打开菜单','按下立即有反馈，动作只执行一次','统一 active focus pending disabled；防重复提交'),
('打开弹窗或侧边面板','先看到面板，再等待内部内容','本地打开状态与请求解耦，慢请求可取消或重试'),
('切换本地标签页','迅速出现对应内容，避免整个页面闪白','状态就近管理；耗时子面板按需加载'),
('跳转常用页面','点击立即显示导航状态，尽早出现可用主体','有限预取、局部 loading、分阶段展示数据'),
('保存或批量操作','知道是否保存，失败后能继续编辑','保留输入和本地草稿，不提前假报云端成功'),
('触屏与键盘操作','手指好点、焦点可见，关闭后回到触发位置','常用触控目标建议至少 44 CSS px；测试 Esc 与 Tab')])
sub('建议的性能验收目标')
p('把按钮按下到可见反馈的 P95 目标设为不高于 100 ms；轻量本地弹窗或标签切换到可用内容的 P95 目标设为不高于 200 ms。重面板和网络完整内容分别制定目标，不能用空壳出现时间替代业务完成时间。以上均是未来验收目标，尚未实测。')
p('标准过渡先控制在约 120 至 200 ms，不让等待动画阻止操作。系统减少动态效果设置和页面级关闭动效选项继续生效。卡顿应先移除主线程阻塞，不能靠拉长动画掩盖。')
sub('为什么计划里保留导航与数据优化')
p('现有侧栏与首页入口关闭自动预取，部分页面主体返回前需要身份或统计数据。即使替换所有按钮，这些等待依然存在。因此将导航反馈、按需预取和局部加载作为换库验收的一部分，同时独立记录其收益，避免把所有改善归功于组件库。')

page();h('风险控制与验收方法')
table(['主要风险','应对措施'],[
('只换外观，使用仍不顺','先对照真实任务演示打开、切换、保存、恢复，再扩展页面'),
('新旧主题与全局样式冲突','统一 token 和层级，限定旧样式作用范围，分批清理'),
('全屏或嵌套浮层出现焦点问题','明确 Portal 容器，测试 Esc 关闭顺序、焦点返回和滚动锁定'),
('迁移影响计时、快捷键或草稿','先包展示层，保留业务数据契约，逐个验证学习流程'),
('依赖和首屏脚本变大','按组件引入，记录路由压缩 JS 体积，不允许无说明的持续增长'),
('源码复制后维护分散','基础组件统一出口，业务页禁止复制第二套基础控件，保留上游版本记录')])
sub('实施时的验证矩阵')
p('至少覆盖 Windows 的 Chromium 浏览器、Android Chrome 和 iOS Safari 的真实设备；桌面鼠标与键盘、手机触控分别验证。逐一检查 light gray dark、减少动效、全屏、手机软键盘、正常网络、慢网络、断网恢复以及已登录和本机模式。缺少某种设备就记录未验证，不能用桌面模拟代替全部真机结论。')
p('在生产构建上固定账号、数据规模和设备，对每个重点动作至少测试 20 次，比较中位数和 P95，同时记录请求耗时与浏览器长任务。20 次用于工程对比；正式 INP 应结合真实访问数据的第 75 百分位观察，良好目标为不高于 200 ms。')
sub('分批发布与回退')
p('在现有 main 上按可独立回退的小批次提交，先替换一个完整操作流程，再推广。保留原业务数据格式，每批次发布前验证草稿和计时记录可恢复。若出现保存异常、关键快捷键失效、全屏浮层不可用等问题，回退对应 UI 批次，并保留已生成的用户数据。')
p('第一批的继续条件是：构建与关键回归通过、三主题一致、常见操作有可感知改善。若仅外观变化或体积成本明显增加而无实际改善，停在第一批重新定位原因，不按原计划机械替换整站。')

page();h('执行建议与资料依据')
p('建议采用本计划，先批准首批范围和体验目标，再按阶段推进。主要界面统一后，可以显著降低“同样的操作在不同页面手感不同”的问题；能改善多少真实响应速度，应由基线和首批对比决定。无需把换 Node.js、换 React 或全站重写作为前置任务。')
sub('规划版本安排')
p('本次只有计划与报告交付，按既有规则把元数据补丁版本从 3.24.3 更新到 3.24.4。实施时引入统一 UI 系统属于新模块，届时按当时 main 的版本升级中间位，不提前占用版本。后续纯修复升级补丁位；只有平台级重构才升级第一位。')
sub('本次仓库核对')
p('开始时成功读取远程 main，本机与 origin/main 提交均为 7a68ddd，提交差异为 0 和 0。工作区原有未提交内容保留；本次仅提交报告、生成源、说明和版本字段。正式迁移尚未启动，计划中的样板、构建和回归均为后续工作。')
sub('官方资料')
refs=[
('shadcn 定位与源码维护','https://ui.shadcn.com/docs'),
('shadcn 的 Next.js 接入','https://ui.shadcn.com/docs/installation/next'),
('Tailwind 版本适配说明','https://ui.shadcn.com/docs/tailwind-v4'),
('Radix 渐进采用与样式职责','https://www.radix-ui.com/primitives/docs/overview/introduction'),
('Radix 服务端渲染','https://www.radix-ui.com/primitives/docs/guides/server-side-rendering'),
('Ant Design 的 Next.js 接入','https://ant.design/docs/react/use-with-next/'),
('MUI 的 Next.js 接入','https://mui.com/material-ui/integrations/nextjs/'),
('Next.js 预取与导航','https://nextjs.org/docs/app/getting-started/linking-and-navigating'),
('INP 的定义和阈值','https://web.dev/articles/inp')]
for title,url in refs: p(title+'  '+url)
sub('源码审阅位置')
p('jackyun-portal/package.json；tailwind.config.ts；src/components/theme-provider.tsx；src/components/layout/sidebar.tsx；src/components/modules/product-card.tsx；src/components/modules/tools/tools-tabs.tsx；src/components/modules/legacy-frame.tsx；src/app/(portal) 下的首页和日程页面。以上为当前本机代码，部分含原有未提交改动。')
for style in D.styles:
 for el in list(style.element.iter(qn('w:color'))): el.attrib.clear();el.set(qn('w:val'),'000000')
 for el in list(style.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
D.save(out/'组件库更换计划与可行性报告3.24.4.docx')
print('saved report')
