"""Build the review DOCX from README.md with the bundled workspace Python."""
from pathlib import Path
import re
import argparse
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
DOCX = ROOT / 'JackYun主站UI现代化实施进展与审查报告.docx'


def font(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def inline(paragraph, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', lambda m: f'{m[1]}（{m[2]}）', text)
    for piece in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text):
        run = paragraph.add_run(piece.strip('`') if piece.startswith('`') else piece.strip('*') if piece.startswith('**') else piece)
        if piece.startswith('`'):
            run.font.name = 'Consolas'
        if piece.startswith('**'):
            run.bold = True


def table(doc, rows):
    count = len(rows[0])
    result = doc.add_table(rows=0, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    widths = {2: [4.1, 12.8], 3: [3.0, 5.0, 8.9], 4: [2.8, 5.0, 4.1, 5.0]}.get(count, [16.9 / count] * count)
    # Version/evidence tables need a short version column and a wide evidence column.
    if count == 3 and rows[0][1] == '实际版本':
        widths = [3.5, 3.0, 10.4]
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    for index, values in enumerate(rows):
        row = result.add_row()
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
        for cell, text, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                item = OxmlElement('w:' + side)
                item.set(qn('w:w'), '95')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            props.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '17365D' if index == 0 else 'F3F6FA' if index % 2 == 0 else 'FFFFFF')
            props.append(shade)
            edges = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    edge.set(qn('w:' + key), value)
                edges.append(edge)
            props.append(edges)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.keep_with_next = index == 0
            inline(p, text)
            for run in p.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string('FFFFFF')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    parser = argparse.ArgumentParser(description='Build a DOCX from the UI modernization Markdown plan or report.')
    parser.add_argument('--input', type=Path, default=ROOT / 'README.md')
    parser.add_argument('--output', type=Path, default=DOCX)
    args = parser.parse_args()
    source = args.input.resolve()
    output = args.output.resolve()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
    font(doc.styles['Normal'], 10.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.18
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12)]:
        font(doc.styles[name], size, name != 'Title')
        if name == 'Title':
            title_ppr = doc.styles[name].element.get_or_add_pPr()
            title_border = title_ppr.find(qn('w:pBdr'))
            if title_border is not None:
                title_ppr.remove(title_border)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(6)
        doc.styles[name].paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.add_run('JackYun 主站 UI 现代化  实施审查')
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string('666666')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    footer.add_run(' 页')
    lines = source.read_text(encoding='utf-8-sig').splitlines()
    index = 0
    while index < len(lines):
        text = lines[index].strip()
        if not text:
            index += 1
            continue
        if text.startswith('>'):
            text = text.lstrip('> ').strip()
        if text.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                index += 1
            table(doc, rows)
            continue
        if text.startswith('#'):
            level = len(text) - len(text.lstrip('#'))
            title = re.sub(r'[、，。；：（）()—–/]', ' ', text.lstrip('#').strip())
            paragraph = doc.add_paragraph(title, 'Title' if level == 1 else 'Heading 1' if level == 2 else 'Heading 2')
            if level == 1:
                paragraph_ppr = paragraph._p.get_or_add_pPr()
                paragraph_border = paragraph_ppr.find(qn('w:pBdr'))
                if paragraph_border is not None:
                    paragraph_ppr.remove(paragraph_border)
        else:
            p = doc.add_paragraph()
            if re.match(r'^\d+\. ', text):
                p.paragraph_format.left_indent = Cm(.6)
                p.paragraph_format.first_line_indent = Cm(-.6)
            inline(p, text)
        index += 1
    doc.core_properties.title = 'JackYun主站UI现代化实施进展与审查报告'
    doc.core_properties.subject = '当前分支实现 实测证据 未完成门禁与审查交接'
    doc.core_properties.author = 'JackYun'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == '__main__':
    main()
