from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name('SEIUE-SCHEDULE-3.44.0-20261025.docx')
doc = Document()
section = doc.sections[0]
section.top_margin, section.bottom_margin = Inches(.72), Inches(.68)
section.left_margin = section.right_margin = Inches(.8)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Subtitle'):
    style = doc.styles[name]
    style.font.name = 'Aptos'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 2'].font.size = Pt(12)
header = section.header.paragraphs[0]
header.text = 'JACKYUN SCHEDULE  |  SEIUE READ-ONLY INTEGRATION'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('实现与安全验收说明 · v3.44.0 · 2026-10-25  |  ')
page = OxmlElement('w:fldSimple')
page.set(qn('w:instr'), 'PAGE')
footer._p.append(page)

def table(headers, rows, widths=None):
    item = doc.add_table(rows=1, cols=len(headers))
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.style = 'Table Grid'
    for i, value in enumerate(headers):
        cell = item.rows[0].cells[i]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), '263B53')
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    item.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row_no, row in enumerate(rows):
        cells = item.add_row().cells
        item.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_no % 2:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'F2F5F8')
                cells[i]._tc.get_or_add_tcPr().append(shade)
    if widths:
        for row in item.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return item

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

title = doc.add_paragraph(style='Title')
title.add_run('SEIUE 与 SCHEDULE 深度同步')
subtitle = doc.add_paragraph(style='Subtitle')
subtitle.add_run('Portal 3.44.0 · 2026 年 10 月 25 日')
doc.add_paragraph('本说明记录 SEIUE 课程、待办和校历融入 JackYun SCHEDULE 的实现方式。课程同步特别处理调课：SEIUE 的个人课程接口返回所选周的实际课程实例，因此成功取得周快照后，日历显示以这份实时数据为准；JackYun 原课表仍保存在本地，不会被同步写入或覆盖。')

doc.add_heading('1. 用户体验与视图合并', level=1)
table(['内容', 'SCHEDULE 中的行为'], [
    ('课程与调课', '日视图时间线与周课表使用 SEIUE 当前周实际返回的课程日期、起止时间、地点和教师。调课后的课程会出现在新日期/课节；成功快照中不再出现的旧课节不残留在该周视图。'),
    ('本地课表保护', 'SEIUE 数据只作为显示快照，不写入本地 dateOverrides、课表导入或云同步记录。未连接、当前周还没有成功快照或主动解除连接时，日程回到 JackYun 本地课表。'),
    ('来源标签', 'SEIUE 单独提供的课显示“SEIUE”；课程名称和起止时间都精确匹配的重复项合并为一个事件，并显示“JackYun + SEIUE”。不通过时间重叠推断同一节课。'),
    ('作业与选课截止', '仅显示源状态明确为 pending 的未完成待办，提供截止、课程组、任务类型、布置人和源端说明（字段可用时）；点击打开学校站点查看或提交，JackYun 不代为提交或勾选完成。'),
    ('校历', '按日期独立展示 SEIUE 校历事件。假期/校历不会转换成课程或课程复习任务。'),
    ('课后复习', '每科设置默认关闭；开启后以 SEIUE 单次课程 ID 生成稳定任务，避免重复创建，并尊重本地完成与取消状态。')
], [1.35, 5.75])

doc.add_heading('2. 周快照与调课规则', level=1)
doc.add_paragraph('SEIUE 个人课程接口的关键字段为 `type=lesson`、`start_time`、`end_time` 和 `custom.id`。研究记录显示 `custom.id` 是周内单次课 ID，跨相邻周不会复用；因此它适合标识具体课次，不是课程模板 ID。课程时间是没有时区偏移的学校墙上时间，按 Asia/Shanghai（UTC+8）严格解析。')
bullet('在今日或周课表视图打开时读取当前所选周；切换查看日期后也会刷新。浏览器标签重新获得焦点或从隐藏状态恢复时重新读取，用户还可主动点击“刷新来源”。')
bullet('只有完整请求并规范化成功的该周快照才会成为显示权威。已连接但请求尚未成功时保留本地课程；失败时保留该账户该周上一次成功的 SEIUE 快照，不清空日程。')
bullet('成功快照被视为该周完整课程列表，所以本地旧日期不会与 SEIUE 新日期并排显示。换周后用相应周的独立快照；本地配置本身始终不变。')
bullet('当前样本尚未验证调课后 `custom.id` 是否变化，也没有确认 adjustment_status 的调课枚举。实现不臆测字段语义，而以所选周实际课程实例的集合及时间为准；源端待机/代课备注只按原值展示。')
bullet('当 SEIUE 确认返回空课程列表时，空列表也是该周有效快照，日历显示没有课程。若上游格式异常或请求失败则不把它当作空表。')

doc.add_heading('3. 连接助手与每日续期', level=1)
table(['步骤', '用户操作与系统反应'], [
    ('安装/更新', '在账户关联页点安装/更新助手，交给 Tampermonkey 或 Violentmonkey；脚本范围固定为 https://chalk-c3.seiue.com/*。'),
    ('获取或续期', '从 JackYun 按钮主动打开已登录的学校页。脚本仅在这一用户发起的一次性流程里观察匹配的个人课程 GET；成功后自动把短期学生连接包复制到剪贴板。普通浏览 SEIUE 页面不会自动捕获。'),
    ('粘贴验证', '返回 JackYun，在连接对话框粘贴连接包。浏览器内存中校验格式后将 Bearer Token 发送到 JackYun API，服务端验证 SEIUE 只读访问、加密保存；前端不把 Token 写进 localStorage、日志或导出。'),
    ('可选每日提示', '开关只保存在本设备。启用时每日第一次打开 SCHEDULE 会提示获取连接信息；用户点击后打开学校页面并自动复制，返回粘贴即可续期。它不自动登录、不后台续期，也不绕过粘贴确认。'),
    ('失败与解除', 'SEIUE 401、授权过期或用户解除连接时服务端删除授权密文并要求重新获取；API 暂时失败时保留已有课程/待办快照。')
], [1.35, 5.75])

doc.add_heading('4. 安全边界', level=1)
bullet('仅支持学生身份和短期 Bearer Token。不采集学校密码、Cookie 或 refresh token；连接包有大小、格式、角色和站点范围校验。')
bullet('浏览器助手只匹配 chalk-c3.seiue.com 页面和 api.seiue.com 的个人课表读取，限制一次捕获，并恢复原 fetch/XHR 调用以不影响学校页面正常请求。脚本没有向 JackYun 自动上传请求；用户粘贴后才通过 JackYun HTTPS API 发送连接数据。')
bullet('服务端上游目标固定为 api.seiue.com 的已验证只读 GET 接口，不接收任意代理 URL。授权包用 AES-256-GCM 加密；没有有效加密密钥时拒绝连接。账户 ID 操作按当前用户隔离。')
bullet('授权保存最长 24 小时；到期或上游 401 后删密文。API 响应使用 private, no-store，只返回规范化且有界的课程、待办和校历字段。禁止在响应、日志、浏览器存储、导出或手表端暴露 Token。')
bullet('无自动提交、源端写入或源端完成状态回写；未经官方机制验证不实现自动续期。用户保留学校账户和剪贴板操作的主动控制权。')

doc.add_heading('5. 验证记录', level=1)
table(['检查', '结果'], [
    ('TypeScript', '通过：tsc --noEmit。'),
    ('SEIUE 定向测试', '29 项通过，覆盖用户隔离、密文与凭据不回传、401/403/过期处理、固定 GET、课程/待办/校历规范化、助手范围、缓存保留、调课到新日期后旧本地位置不显示、精确重复合并及不同课程时间重叠不误删。'),
    ('生产构建', '通过：next build --webpack 成功完成优化编译、TypeScript 检查、98 个静态页面生成和构建 trace 收集；/schedule、SEIUE 账户与快照 API 均出现在路由清单。'),
    ('浏览器真实账户验收', '需在受保护预览中由用户登录并重新关联短期 Token。未完成此步骤前，不声称真实学校账户链路已验收。'),
    ('部署', '合并后核验 Vercel 生产部署完成，并检查线上版本和 SCHEDULE 路由；PR 或预览成功不等于生产部署成功。')
], [2.0, 5.1])

doc.add_heading('6. 发布记录', level=1)
doc.add_paragraph('本次属于新功能，Portal 版本从 3.43.2 提升到 3.44.0。源代码位于独立功能分支，合并前要求定向测试、类型检查、生产构建和安全边界检查通过。')
doc.add_paragraph('真实账户验收不得通过聊天发送 Token、Cookie、密码或包含 Authorization 请求头的 HAR/截图。遇到需要重新连接的提示时，请只在预览页面粘贴脚本生成的临时连接包。')

doc.save(OUT)
print(OUT)
