from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).with_name("SEIUE-Schedule-3.42.9-20261024.docx")
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.78)
sec.right_margin = Inches(.78)

styles = doc.styles
for name in ("Normal", "Title", "Heading 1", "Heading 2", "Subtitle"):
    st = styles[name]
    st.font.name = "Aptos"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.color.rgb = RGBColor(26, 54, 93)
styles["Heading 2"].font.size = Pt(12)
styles["Heading 2"].font.color.rgb = RGBColor(40, 90, 130)

header = sec.header.paragraphs[0]
header.text = "JACKYUN PORTAL  |  SEIUE × SCHEDULE"
header.style = "Caption"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("内部实现说明 · 3.42.9 · 2026-10-24    |    ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Light Shading Accent 1"
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = text
        shade(c, "1A365D")
        for run in c.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def numbered(text):
    doc.add_paragraph(text, style="List Number")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("版本整合与安全验收记录\n").bold = True
title = doc.add_paragraph(style="Title")
title.add_run("SEIUE SCHEDULE\n只读同步与 Main 同步恢复")
sub = doc.add_paragraph(style="Subtitle")
sub.add_run("功能分支基线：3.42.8  →  3.42.9    |    记录日期：2026 年 10 月 24 日")
doc.add_paragraph("本说明记录把 Main 的账户状态读取恢复与安全诊断整合进 SEIUE 集成分支后的实现边界、已验证内容和仍需的验收步骤。外部账户仍为只读；本地 SCHEDULE 课表及尚未同步的本机操作继续由原同步队列管理。")

doc.add_heading("1. 变更摘要", level=1)
table(["领域", "行为"], [
    ("Main 只读恢复", "对 Supabase REST 的 GET/HEAD 读取，仅对网络/超时及 408、429、5xx 等可恢复失败进行有界恢复；明确 Retry-After 时遵守退避，不重试用户取消的请求。"),
    ("同步写入", "不对写入请求做自动重放；本地持久队列、确认与冲突选择语义保持原样，避免重复应用或覆盖并发编辑。"),
    ("会话权威状态", "同一已验证身份与会话的并发检查可合并；已知有效会话只在短暂上游 503 场景提供短暂连续性，权限拒绝和过期会话仍 fail-closed。"),
    ("失败诊断", "新增经同源、登录态、请求大小与速率限制的诊断入口；只接受白名单字段，不发送内容、错误消息/堆栈、token 或账户凭据。"),
    ("SEIUE 快照", "课程、待办及校历继续使用固定官方域名的只读请求；规范化响应不包含授权材料，失败保留上一份成功快照。")
], [1.55, 5.5])

doc.add_heading("2. 同步安全边界", level=1)
doc.add_heading("2.1 本机操作与云端写入", level=2)
bullet("读取恢复只作用于 GET/HEAD 数据读取，且限定到 Supabase REST 路径；不会自动重试 Azure 写入或本地队列中的变更操作。")
bullet("诊断上报与业务同步分离。诊断仅包含已枚举的阶段、状态码、耗时和安全错误类别，并做载荷字节数、同源及用户级节流检查。")
bullet("同步初始化或读取失败时，不清理 IndexedDB/local workspace 的待提交记录；冲突仍需用户明确选择。")
bullet("SEIUE 与日程采用叠加视图，不将课程回写为覆盖本地 schedule 的 dateOverrides。外部授权到期或解除连接时删除授权密文。")

doc.add_heading("2.2 凭据与个人数据", level=2)
bullet("SEIUE token 仅由用户在设置页粘贴或使用受限的浏览器连接助手复制；连接助手在学校页面内短暂观察一个符合条件的个人课程请求，凭据只存在页面内存并需用户点击复制。")
bullet("服务端固定请求 api.seiue.com 的已验证 HTTPS GET 端点，不接受调用方提供代理 URL；响应只返回规范化课程、待办与校历字段。")
bullet("连接信息使用服务端 AES-256-GCM 密钥加密，最长 24 小时；不收集学校密码、Cookie 或 refresh token。无加密密钥时拒绝保存。")
bullet("不得在日志、错误正文、浏览器存储、导出或诊断数据中记录 token、Cookie、学校密码或完整上游响应。")

doc.add_heading("3. SEIUE 与 SCHEDULE 行为", level=1)
table(["数据", "同步规则", "明确限制"], [
    ("课程", "按用户选择周读取，以 SEIUE 实际起止时间为准；无时区时间按 Asia/Shanghai 严格解析。以 custom.id 作为单次课源 ID。", "作为外部事件叠加；课程复习任务逐科开关，默认关闭，并使用稳定源 ID 防止重复。"),
    ("作业/资料", "分页读取；以 todo id 去重；使用顶层 status 作为完成状态，截止时间从对应 related 字段映射。", "提供 SEIUE 首页新标签跳转；JackYun 不代提交，也不向 SEIUE 写完成状态。"),
    ("选课截止", "从已知选课待办类型识别，按同一待办状态和关联截止字段映射。", "不会与普通作业混淆；未知类型不猜测映射。"),
    ("校历", "作为独立校历事件展示；只接受已知结构并规范化。", "校历事件不会混入课程或生成课程复习。")
], [1.1, 3.3, 2.65])

doc.add_heading("4. 回归验证", level=1)
table(["检查", "结果", "覆盖范围"], [
    ("SEIUE + 同步诊断/恢复定向测试", "通过：49/49", "账户隔离、token 不回传、401/403/过期、固定只读上游、课程/待办/校历映射、按周缓存、连接助手约束、失败诊断过滤、离线队列及冲突保留。"),
    ("Azure 权威会话测试", "通过：7/7", "用户身份、权限拒绝、短暂故障、会话缓存边界及原 JWT 到期约束。"),
    ("Azure Functions 适配器测试", "未运行", "当前工作区 azure/backend/node_modules 未安装锁定依赖 @azure/functions@4.16.5；不是测试失败，需在可安装依赖的 CI/构建环境运行。"),
    ("Next.js webpack production build", "通过", "next build --webpack 已成功完成 TypeScript 检查、页面生成与 trace 收集；/schedule 和 SEIUE 连接及快照 API 均进入产物路由清单。"),
    ("浏览器真实账户与云端队列验收", "待完成", "用户需在最新预览重新确认账户和本机同步状态；此前页面存在待上传操作，不能用自动点击重试代替核对。")
], [1.65, 1.2, 4.2])

doc.add_heading("5. 已知恢复边界与验收动作", level=1)
doc.add_paragraph("本次恢复针对 Main 中确认的短暂账号状态读取失败与安全诊断能力。浏览器曾出现过同步请求 400 INVALID_OPERATIONS 及未完成的本机操作；这类确定性请求校验错误不是短暂网络读取问题，因此不能通过重复发送解决。实现保留本地队列和服务端数据，不会擅自丢弃、拆分或重放该批操作。")
numbered("先等待预览部署指向本次合并后的提交，并核对页面显示版本为 3.42.9。")
numbered("打开现有日程页时先检查同步状态及待处理项目数量；不要在未确认冲突明细前点“强制同步/清理本地数据”。")
numbered("如果出现待处理本地更改，先从同步中心查看每一项的本地与云端版本，再按用户选择解决；同步写入只在用户确认后进行。")
numbered("SEIUE 短期 token 过期时，通过设置页自动连接助手重新连接；不要把 token、Cookie、HAR 或包含 Authorization 的截图发到聊天或提交到仓库。")
numbered("完成后再核对课程周、作业跳转、回站刷新、解除连接、离线时保留旧快照等流程。")

doc.add_heading("6. 发布与版本记录", level=1)
bullet("功能分支：feat/seiue-schedule-sync；Main 当前基线为 4c55462c1e1b1309826f5ed4b69fcf36881cfccf。")
bullet("应用版本统一更新为 3.42.9：package.json、package-lock.json、APP_VERSION 与更新历史同步。")
bullet("本记录用于提交到仓库 docs/ 目录。完成分支构建与人工验收前，不把预览通过等同于生产发布，也不宣称已经合并 Main 或生产成功。")

doc.add_heading("7. 本次实施范围", level=1)
doc.add_paragraph("本次工作是将 Main 上的有界只读恢复及失败诊断变更整合到现有 SEIUE 深度同步分支，目的在于改善会话读取瞬时失败时的 SCHEDULE 使用连续性，同时保留用户本机待同步变更。自动续期、SEIUE 写入、代提交与删除/覆盖 SCHEDULE 本地事件都不在启用范围。")

doc.save(OUT)
print(OUT)
