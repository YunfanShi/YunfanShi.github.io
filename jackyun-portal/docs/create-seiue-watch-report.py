from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name('SEIUE-Schedule-Watch-3.43.0-20261024.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.68)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)

for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Subtitle'):
    style = doc.styles[name]
    style.font.name = 'Aptos'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 2'].font.size = Pt(12)
for element in (doc.styles['Title']._element,):
    properties = element.find(qn('w:pPr'))
    if properties is not None:
        borders = properties.find(qn('w:pBdr'))
        if borders is not None:
            properties.remove(borders)

header = section.header.paragraphs[0]
header.text = 'JACKYUN PORTAL AND WATCH  |  SEIUE READ-ONLY SYNC'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('实现与安全验收说明 · v3.43.0 · 2026-10-24  |  ')
page = OxmlElement('w:fldSimple')
page.set(qn('w:instr'), 'PAGE')
footer._p.append(page)

def table(headers, rows, widths=None):
    item = doc.add_table(rows=1, cols=len(headers))
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.style = 'Table Grid'
    for i, value in enumerate(headers):
        cell = item.rows[0].cells[i]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '263B53')
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    header_props = item.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement('w:tblHeader')
    repeat_header.set(qn('w:val'), 'true')
    header_props.append(repeat_header)
    keep_header = OxmlElement('w:cantSplit')
    header_props.append(keep_header)
    for n, row in enumerate(rows):
        cells = item.add_row().cells
        row_props = item.rows[-1]._tr.get_or_add_trPr()
        row_props.append(OxmlElement('w:cantSplit'))
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n % 2:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F2F5F8')
                cells[i]._tc.get_or_add_tcPr().append(shading)
    if widths:
        for row in item.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return item

def bullet(value):
    doc.add_paragraph(value, style='List Bullet')

title = doc.add_paragraph(style='Title')
title.add_run('SEIUE 日程同步与手表端接入说明')
subtitle = doc.add_paragraph(style='Subtitle')
subtitle.add_run('Portal 3.43.0 · Watch 3.1.0-beta.1 · 2026 年 10 月 24 日')
doc.add_paragraph('这份说明记录 SEIUE 与 JackYun SCHEDULE 的只读同步改进，以及 Watch 日程里的待办呈现。待办只在源状态明确为未完成时显示，并附带截止时间和可用的课程上下文。SEIUE Token 留在 Portal 服务端加密区，不发送给 Watch；任务仍由用户回到 SEIUE 查看或提交。')

doc.add_heading('1. 本次变更', level=1)
table(['位置', '实现'], [
    ('SCHEDULE 作业与选课卡片', '只显示顶层 status 为 pending 的作业、资料与选课待办。已完成、已忽略、已失效及未知状态不会被标成待办。'),
    ('任务必要信息', '校验 Asia/Shanghai 本地截止时间并统一显示到分钟；展示课程组、作业/资料类型、布置人、布置时间和任务说明（源端有提供时）。'),
    ('Watch 日程', '服务端读取当前 JackYun 用户已关联的 SEIUE 只读快照，将近 14 天内及无截止时间的未完成任务提供给 Watch 日程；列表最多 200 项。'),
    ('Watch 截止提醒', '遵循 Watch 原有“任务截止”通知开关；提示仅对状态为 pending 且有合法截止时间的 SEIUE 待办建立，不进入可写任务完成队列。'),
    ('离线与上游故障', 'Watch API 失败时通过现有模块快照恢复保留上次成功的 SEIUE 日程；SEIUE 失败不清除 JackYun 本地课表或待同步队列。')
], [1.55, 5.55])

doc.add_heading('2. SEIUE 字段映射', level=1)
table(['界面字段', '源字段和规则'], [
    ('待办 ID', '顶层 id；作为稳定源 ID 使用。'),
    ('是否显示', '仅接受顶层 status=pending。related.status 表示作业流转状态，不用于判断待办完成。'),
    ('截止时间', '作业与资料使用 related.expired_at；选课截止使用 related.end_at。非法日期不转成时间提醒。'),
    ('任务上下文', '课程组 related.group.name；布置人 related.creator.name；作业或资料类型来自已知 related.labels.type；描述从 related.content 的 Draft.js blocks[].text 读取。'),
    ('分页与范围', '待办请求仍分页读取，使用固定 api.seiue.com HTTPS GET 路径；只请求 status_in=pending。Watch 只输出未来 14 天内、逾期未完成及无截止时间的项目。')
], [1.45, 5.65])

doc.add_heading('3. Watch 端只读边界', level=1)
bullet('Watch API 使用已配对设备的用户身份，再按 user_id 和 provider=seiue 读取关联，不接受客户端指定账户 ID，也不向设备返回密文或 Token。')
bullet('仅对授权仍有效的 SEIUE 账户解密 AES-256-GCM 密文并请求固定官方只读接口。过期或上游 401 时删除该用户对应的授权密文并要求重新连接；403 不擦除仍有效的密文。')
bullet('SEIUE 待办仅进入 agenda/seiueAgenda 呈现和截止提醒。任务详情没有“标记完成”操作，Watch 的任务写入队列不能修改 SEIUE 状态。')
bullet('响应设置 private, no-store。账户查询、密文查询、过期清除和状态更新均以当前用户 ID 限定。失败时不记录原始上游响应、任务内容或授权信息。')

doc.add_heading('4. 凭据与个人数据保护', level=1)
bullet('连接只接收学生角色的短期 Bearer Token、Reflection ID、School ID 和 SEIUE 学校前端网址；不收集学校密码、Cookie 或 refresh token。')
bullet('Token 在服务端 AES-256-GCM 加密保存，最长 24 小时；没有配置有效加密密钥时拒绝关联。到期或 SEIUE 返回 401 时删除密文。')
bullet('Watch 使用 JackYun 设备凭据访问自身快照 API。SEIUE Token 不写入 Watch 响应、APK、设备本地设置、提醒内容、日志或导出数据。')
bullet('自动续期、SEIUE 写入、代提交和源端完成状态回写仍保持关闭。用户点击学校链接后由 SEIUE 官方页面完成查看或提交。')

doc.add_heading('5. 验证清单与结果', level=1)
table(['检查', '结果与边界'], [
    ('SEIUE、Watch 映射和路由定向测试', '47/47 通过：截止时间校验、课程/类型/布置人描述映射、明确待办过滤、分页、认证隔离、Watch 用户隔离、过期密文清理和只读端点。'),
    ('Portal TypeScript 检查', '通过：tsc --noEmit。'),
    ('Watch Android 单元测试、Lint、Release 构建', '本地通过：:app:testReleaseUnitTest、:app:lintRelease、:app:assembleRelease。产物为 1,482,571 字节，无 native library，低于 4 MB 限制。lint 报告生成；目标设备仍需确认腕上版式和通知送达。'),
    ('在线 SEIUE 账户验收', '需在最新预览登录并重新关联短期 Token；本说明不包含任何生产凭据或真实响应样本。'),
    ('实体手表验收', '无已连接设备证据时不宣称通知送达和腕上版式已通过；需在目标设备确认日程显示、提醒权限与截止通知。')
], [2.25, 4.85])

doc.add_heading('6. 版本与发布边界', level=1)
doc.add_paragraph('Portal 功能版本从 3.42.10 升为 3.43.0；Watch 从 3.0.0-beta.1 升为 3.1.0-beta.1，versionCode 从 21 升为 22。两个项目位于不同 Git 仓库，Portal PR 与 Watch 源码/CI 独立验证。PR 合并或预览可用不等于生产部署成功；生产发布后仍须核对 Vercel 部署状态、线上版本和日程路由。')
doc.add_paragraph('由于 SEIUE 连接令牌为短期凭据，真实账户端到端验证必须在预览网页内由账户持有人完成。不要把 Token、Cookie、完整 HAR 或包含 Authorization 请求头的截图发到聊天或提交仓库。')

doc.save(OUT)
print(OUT)
