"""Create the implementation report using the bundled Python runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.85)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.add_paragraph('JackYun 实时课程时间线实施说明', 'Title')
doc.add_paragraph('版本 3.10.0    日期 2026 年 10 月 12 日')
doc.add_paragraph('本次把今日课表升级为实时课程显示，让使用者直接看到当前课程、持续时间和距离下课的时间。沿用已有课表数据与卡片风格，新增功能按版本规则从 3.9.1 升至 3.10.0。')
for title, body in [
    ('当前课程显示', '课表卡片顶部优先显示正在上的课程全名、开始和结束时间、总时长、已上分钟、下课倒计时及进度条。下方保留完整课程列表、课程颜色、教室和教师；每一节均显示时段与分钟数。当前课仅增加文字标记，保留原有浅色卡片和分隔线。'),
    ('课间与时间边界', '倒计时逐秒刷新，按当前系统时间与结束时刻之差计算。上课开始时刻包含在当前课中，下课结束时刻不包含，边界自动切换。课间显示下一节及距上课时间；尚未上课和今日课程结束分别提示。相邻课程无间隙时直接切换下一课，超过一小时使用小时分钟秒显示。'),
    ('日期和模式', '课表日期与时钟沿用浏览器本地时间，与已有日期选择保持一致。查看其他日期时标题显示当日课表，不宣称正在上课或显示实时倒计时；跨午夜后旧日期也按历史日期处理。假日模式保留暂停说明，不显示课程倒计时。课表周次与单日调课继续使用现有 sessionsForDate 规则。'),
    ('实现和性能', '新增独立 ClassTimeline 客户端组件与纯时间计算函数，替换原今日课表块。逐秒时钟只刷新这一小块，不让整个学习控制台逐秒重渲染；隐藏页面跳过刷新，重新聚焦或恢复可见时立即取当前时间。返回页面按时钟重新计算，不累计暂停前的秒数。清理定时器与事件监听，倒计时不每秒向读屏程序播报。'),
    ('参考项目与数据', '参考 ClassIsland 官方项目展示当天课表和当前课程的思路，并按本次需求实现网页倒计时。参考地址为 https://github.com/ClassIsland/ClassIsland 。未复制其代码或接入其服务。此次只读现有本机课表及同步结果，不写入学习状态、草稿或日程；无数据库迁移。异常或非正时长时段不进入实时计算。'),
]:
    doc.add_heading(title, 1); doc.add_paragraph(body)
doc.add_heading('验证和交付', 1).paragraph_format.page_break_before = True
doc.add_paragraph('新增 3 项边界规则测试通过，覆盖上下课精确时刻、相邻课程、长课程、午夜、其他日期、空课表、无效时段和暂停后返回。186 项学习回归通过。手机 390 像素及桌面 1440 像素浏览器验证逐秒变化、进度、课间切换、下一节、课程结束、日期导航、午夜与假日模式。生产构建含 TypeScript 检查通过，lint 无错误且仅有两项既有警告。')
doc.add_paragraph('完整验证运行 package.json verify 中各项脚本并继续执行失败项后续检查，共 555 项 Node 测试，550 项通过，5 项既有静态断言失败。legacy-ui 4 项、reader-admin 1 项的失败延续既有记录；check:companion 仍因缺少 unzip 无法完成。私钥扫描及其余检查通过。系统 npm 入口缺失时直接使用 node 执行仓库同名脚本，未因此省略验证。')
doc.add_paragraph('开始时本机与远端 main 同为 66fe0aa，版本 3.9.1，工作区干净。清单、锁文件、APP_VERSION、README 与更新历史统一更新 3.10.0，代码与本说明提交 main 并推送，结束后核对提交一致性。需要回退时撤销该功能提交，保留原有课表数据。截图来自隔离测试，未操作个人生产账号；真实跨设备与手机真机未验收。')
doc.add_heading('手机界面截图', 1)
doc.add_paragraph('测试时刻为 09 时 22 分 30 秒，数学课 09 时开始、09 时 45 分结束，剩余 22 分 30 秒，课程进度为 50%。')
doc.add_picture('artifacts/verification/cue-roadmap/timeline-390.png', width=Inches(2.7))
footer = section.footer.paragraphs[0]
footer.add_run('JackYun 3.10.0    ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
output = Path('docs/JackYun-3.10.0-Class-Timeline-20261012.docx')
doc.save(output); print(output.resolve())
