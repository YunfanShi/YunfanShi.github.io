"""Release explanation for the dashboard layout correction."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.85)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.add_paragraph('JackYun 首页 CUE 布局修复', 'Title')
doc.add_paragraph('版本 3.9.1  |  2026 年 10 月 12 日')
for title, body in [
    ('问题与修复目标', '3.9.0 把 CUE 大看板放在首页问候之前，打断了原有首页顺序；独立数字卡片与重复入口也使页面显得拥挤。按本次反馈，CUE 应属于“接下来学什么”，并沿用原有首页风格。此次按布局问题修复递增补丁版本至 3.9.1。'),
    ('页面位置与视觉处理', '首页依次保留问候、JackYun AI、学习概览和快速收集，随后展示“接下来学什么”。CUE 嵌入这一块原有卡片，再接云端学习任务。取消 CUE 独立外框、大号数字背景卡片和重复操作按钮；使用紧凑文字统计、轻分隔线及原有蓝色文字链接。手机统计以三列排布，较宽屏幕五列排列。'),
    ('入口与交互', '每科保留一个“开始学习”文字入口，点击先进入当前学科预览，再确认开始。顶部“查看全部与时间预算”进入完整待学预览。已有草稿继续显示原学科或联合学习名称及已评价进度，继续入口恢复原会话；详细统计口径收进可展开的说明。云端任务管理与专注计时入口保留。'),
    ('规则与实现范围', '仅调整 dashboard/page.tsx 的组件位置和 CueDashboard 的布局，不改 CUE 到期、自评、冷却、预算或同步规则。今日待学、已学、未学／修改、薄弱和冷却数量仍按既有逻辑计算；分类重叠仍有明确说明。没有新增数据库迁移，也不重新生成用户草稿或日程。清单、锁文件、APP_VERSION、README 和站内更新记录同步为 3.9.1。'),
    ('验证与发布限制', '已运行手机 390 像素、桌面 1440 像素浏览器回归，确认 CUE 位于“接下来学什么”内部且问候在前，继续验证学科预览、预算确认、草稿恢复、读取失败和完成反馈。lint 无错误，存在两项既有未使用变量警告；生产构建含 TypeScript 检查通过。此补丁未重跑全仓库 verify；3.9.0 已记录的五项旧静态断言失败及缺少 unzip 的扩展检查限制仍存在。真实登录账号与手机真机未在本次验收。'),
    ('仓库核对与回退', '开始时本机与远端 main 同为 4162eeb（3.9.0），工作区干净。修复直接提交并推送 main，不建立新分支，完成后再次核对远端提交及工作区。需要撤销时使用该布局修复提交的 git revert；文档、代码与版本清单随提交一起管理。'),
]:
    if title == '仓库核对与回退':
        doc.add_page_break()
    doc.add_heading(title, 1)
    doc.add_paragraph(body)
doc.add_heading('最终界面验证', 1)
doc.add_paragraph('以下截图来自隔离测试数据的手机视口，展示 CUE 与原有“接下来学什么”卡片的整合；数据不代表个人账号。')
doc.add_picture('artifacts/verification/cue-roadmap/dashboard-390.png', width=Inches(3.2))
doc.add_paragraph('浏览器回归：npm run test:learning-roadmap-ui；布局截图在生产构建后重新核对。构建：npm run build。静态检查：npm run lint。')
footer = section.footer.paragraphs[0]
footer.add_run('JackYun 3.9.1  ·  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
output = Path('docs/JackYun-3.9.1-Dashboard-CUE-Layout-20261012.docx')
doc.save(output)
print(output.resolve())
