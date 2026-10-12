"""Generate release notes with the Codex bundled Python and python-docx."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

STAGES = {
 '3.7.0': ('CUE 首页看板', [
  ('改动与使用流程', '首页第一块改为今天的 CUE，展示今日待学、今日已学、未学或已修改、薄弱和熟练冷却中的数量。各学科按待学数量降序显示，查看待学和开始学习均先打开预览；预览确认后才创建学习会话。已有同范围草稿优先继续，保持原顺序、评价与作答。AI 介绍及工具入口移到后面。'),
  ('数据口径', '统计按 CUE ID 去重，复用现有到期判断和熟练冷却规则。分类存在重叠，不应相加。当天学过但内容随后修改的 CUE 重新待学，未修改的当天学习内容不重复出现。熟练冷却到期当日恢复可学。今日已学计算当前内容已检索的 CUE 数量，不能解释为永久掌握量。'),
  ('实现与兼容', '共享只读读取器保留已有规范化、旧数据迁移和学科对齐行为；读取损坏数据时明确报错并保留原始存储，初始化失败不会写入空数据。首页通过本机存储、账号同步应用事件、其他标签页变化、窗口重新聚焦及每分钟日期检查刷新。未添加数据库表，沿用现有同步协议。'),
  ('入口参数', '/schedule 支持 view=cue、subject、unit、filter 和 draft。filter 默认 due，薄弱练习使用 weak。参数通过 URLSearchParams 编码，进入后使用当前数据校验范围；失效草稿和已删除内容会提示。链接本身不执行日程写入。'),
 ], '10 项看板及熟练规则单元测试通过；TypeScript 检查通过；390 和 1440 像素浏览器验证通过，确认预览不建草稿、明确开始才保存会话。', 'dashboard-390.png'),
 '3.8.0': ('复习反馈与学习分析', [
  ('错题分批复习', '今日到期总数独立统计，每批最多 50 题。批次结束后重新读取权威到期队列，展示本批完成和加载下一批；只有权威数据确认没有剩余待复习内容时才显示今日完成。失败允许重试，不丢失本机作答。'),
  ('离线反馈', '本机复习结果先持久保存，离线操作沿用账号专属队列。待同步期间显示本机已记录而不是宣称云端全部完成。统计随归档和评价更新；恢复网络后由已有离线同步服务提交并重新生成快照。'),
  ('学习结束反馈', '学习完成后展示本次 CUE 和题目数量、稳定与部分及未掌握自评数量、剩余待学和下一次复习日期，提供继续下一组和返回首页。记录保存失败时保留草稿并显示错误，不进入成功反馈。'),
  ('表现统计', '增加近 7 天、近 30 天和全部历史范围，默认近 7 天；用账号学习日期按日筛选证据，标明样本数。自评稳定比例不解释为客观正确率。学科和单元提供薄弱练习入口，先查看当前可练习内容再确认开始，AI 报告不修改状态或安排日程。'),
 ], '7 项错题分页、离线完成和统计边界测试通过，186 项学习及 112 项同步回归通过。手机与桌面表现图和薄弱入口浏览器检查通过；实际 ReviewSession 与 OfflineWorkspaceSync 组件在隔离 API 响应下验证 0、50、51、120 题、失败重试、IndexedDB 离线队列及联网恢复。学习完成浏览器检查确认记录持久保存、草稿移除与首页返回。真实账号跨设备和手机真机仍需独立验收。', 'insights-390.png'),
 '3.9.0': ('每日学习预算', [
  ('预算选择', '提供 15、30、60 分钟三个预算，默认 30 分钟。使用现有每条 CUE 预计 4 分钟的估时形成学习预览，不表示实际用时保证。预算计算只生成推荐范围，不创建会话或任务。'),
  ('排序与剩余内容', '先安排已复习且到期的 CUE，再安排未学或修改后待学内容；同组内按薄弱程度优先。遵守已有当日复习和熟练冷却规则，按 ID 去重。超出预算内容明确显示并允许逐条手动加入，实际预计时长随选择更新。'),
  ('确认与日程', '已有草稿优先恢复。开始学习需明确确认；加入日程使用当前范围生成一项手动学习任务，携带准确 CUE ID，防止重复点击生成等价任务。取消预览不写入日程。任务开始时仍校验内容是否存在和是否处于冷却。'),
  ('发布及回退', '三个阶段都沿用现有学习与同步能力，不重构平台。版本依次为 3.7.0、3.8.0、3.9.0，并同步更新清单、锁文件、站内版本、更新记录和 README。回退使用相应提交的 git revert，保留用户 CUE、评价和草稿数据，避免恢复旧数据覆盖新记录。'),
 ], '新增路线图规则测试 15 项全部通过，学习领域 186 项、同步 112 项及 PWA 22 项全部通过。生产模式手机和桌面浏览器验证了预算切换、手动超预算加选、精确范围开始、明确加入日程、重复点击去重、草稿优先、损坏数据保留及键盘焦点。实际复习组件覆盖 0、50、51、120 题、下一批读取失败重试、离线空缓存及真实 IndexedDB 队列联网恢复。完整检查共运行 552 项 Node 测试：547 项通过，5 项既有静态断言失败；并非完整检查全部通过。生产构建、TypeScript、零错误 lint 与敏感信息检查通过。', 'budget-390.png'),
}

def build(version):
 title, sections, evidence, screenshot = STAGES[version]
 doc = Document()
 for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
 sec = doc.sections[0]
 sec.page_width = Cm(21); sec.page_height = Cm(29.7)
 sec.top_margin = sec.bottom_margin = Cm(2); sec.left_margin = sec.right_margin = Cm(2.2)
 for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
  style = doc.styles[name]; style.font.name = 'Microsoft YaHei'; style.font.color.rgb = RGBColor(0, 0, 0)
  style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
 doc.styles['Normal'].font.size = Pt(11); doc.styles['Normal'].paragraph_format.space_after = Pt(8)
 doc.styles['Normal'].paragraph_format.line_spacing = 1.2
 doc.styles['Title'].font.size = Pt(24); doc.styles['Heading 1'].font.size = Pt(15)
 if version == '3.9.0':
  doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
  doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
  doc.styles['Normal'].paragraph_format.line_spacing = 1.15
 doc.add_paragraph(f'JackYun {title}实施说明', 'Title')
 doc.add_paragraph(f'版本 {version}    日期 2026 年 10 月 12 日')
 doc.add_paragraph('本说明面向项目使用者与后续维护者，记录学习体验路线图的具体实现、统计含义、操作流程、验证证据及发布限制。核心目标是从首页知道今天要学什么，确认后快速进入现有学习流程。')
 for heading, body in sections:
  doc.add_heading(heading, 1); doc.add_paragraph(body)
 doc.add_page_break(); doc.add_heading('验证与验收', 1); doc.add_paragraph(evidence)
 doc.add_heading('重点场景', 1)
 doc.add_paragraph('检查零数据、多学科、分类重叠、当日修改、熟练到期、午夜换日以及草稿对应内容删除；覆盖手机与桌面布局、键盘入口和页面运行错误。错题队列应覆盖 0、50、51、120 题，离线提交、重新联网和读取失败分别确认。')
 doc.add_heading('测试环境与限制', 1)
 doc.add_paragraph('浏览器使用本机 Microsoft Edge 的无头模式，390 像素模拟手机和 1440 像素模拟桌面。学习数据使用隔离测试上下文，不操作个人生产记录。本机截图不是手机真机或真实账号跨设备验收；账号权限与外部服务仍需独立验证。')
 doc.add_heading('证据与交付', 1)
 evidence_path = Path('artifacts/verification/cue-roadmap/final-results.txt')
 if evidence_path.exists(): doc.add_paragraph(evidence_path.read_text(encoding='utf-8'))
 else: doc.add_paragraph('开始时本机 HEAD 与远端 main 均为 7a5705515853c3aa5c02573ee921bcbd616b3766，工作区干净。测试命令与截图保存在本机验证目录；后续阶段另附版本说明。')
 image = Path('artifacts/verification/cue-roadmap') / screenshot
 if image.exists():
  doc.add_page_break(); doc.add_heading('界面验证截图', 1)
  doc.add_paragraph('以下为隔离测试数据下的手机视口，用于核对布局与交互入口。')
  doc.add_picture(str(image), width=Cm(10 if version == '3.9.0' else 7.2))
 footer = sec.footer.paragraphs[0]; footer.add_run(f'JackYun {version}   ')
 field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
 out = Path('docs') / f'JackYun-{version}-Learning-Roadmap-20261012.docx'
 doc.save(out); print(out.resolve())

if __name__ == '__main__':
 for version in sys.argv[1:] or STAGES: build(version)
